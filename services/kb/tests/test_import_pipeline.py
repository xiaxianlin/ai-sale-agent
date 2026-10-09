# 知识产权归甲方所有（COM-IP-01）——AI 外贸销售智能体系统交付物，许可与再分发依双方约定。
"""导入流水线单元测试（FR-KB-021 / SR-KB-021 §9：三段校验/部分成功/查重/
公式注入/分片续跑/瞬时降级/异步状态机/留痕）。

本文件同时导出共享测试桩（StubField / StubSchemaProvider / RecordingHook /
build_xlsx_bytes / build_csv_bytes），供 test_import_api.py 复用。
"""
from __future__ import annotations

import asyncio
import csv
import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from services.kb.core.import_service.pipeline import (
    ERROR_BIZ_RULE,
    ERROR_DUPLICATE,
    ERROR_INTERNAL,
    ERROR_SCHEMA_VIOLATION,
    STATUS_DONE,
    STATUS_FAILED,
    STATUS_QUEUED,
    STATUS_RUNNING,
    ImportConfig,
    ImportService,
    ImportValidationError,
    RowIssue,
    TransientImportError,
    build_error_report_xlsx,
    dedup_fingerprint,
    workbook_to_bytes,
)


# --------------------------------------------------------------------------- 共享测试桩


@dataclass(frozen=True)
class StubField:
    """注册表字段桩（与 FR-KB-020 注册表鸭子同形）。"""

    key: str
    label: str
    required: bool = False
    options: tuple[str, ...] = ()
    help: str = ""
    example: Any = ""
    widget: str = ""


class StubSchemaProvider:
    """Schema 注册表桩：product / customer 两库。"""

    def __init__(self) -> None:
        self._fields: dict[str, tuple[StubField, ...]] = {
            "product": (
                StubField("sku", "SKU", required=True, help="产品唯一货号", example="SKU-001"),
                StubField("category", "产品大类", options=("电子", "家居", "服装"), example="电子"),
                StubField("price", "参考单价", help="美元", example="12.50"),
            ),
            "customer": (
                StubField("company", "公司名称", required=True, example="Acme Trading"),
            ),
        }

    def fields(self, kb_type: str) -> tuple[StubField, ...]:
        try:
            return self._fields[kb_type]
        except KeyError:
            raise KeyError(f"未注册的知识库类型：{kb_type}") from None

    def dedup_fields(self, kb_type: str) -> tuple[str, ...]:
        """product 库查重追加 sku 键（库特定查重口径示例）。"""
        return ("sku",) if kb_type == "product" else ()


@dataclass
class RecordingHook:
    """条目创建钩子桩：记录创建动作并派发递增 id。

    - fail_rows：命中标题的行**持续**失败（模拟该行写入故障）；
    - transient=True 抛 TransientImportError（分片重试 + internal 降级路径）；
      False 抛 RuntimeError（任务级 failed → 水位续跑路径）。
    """

    items: list[Any] = field(default_factory=list)
    fail_rows: set[str] = field(default_factory=set)
    transient: bool = True

    def __call__(self, item: Any) -> int:
        if item.title in self.fail_rows:
            if self.transient:
                raise TransientImportError("模拟存储抖动")
            raise RuntimeError("模拟存储不可用")
        self.items.append(item)
        return len(self.items)  # 递增条目 id：1, 2, 3, ...


def build_xlsx_bytes(header: Sequence[str], rows: Sequence[Sequence[Any]]) -> bytes:
    """在内存构造 xlsx 上传件（表头 + 数据行）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    return workbook_to_bytes(workbook)


def build_csv_bytes(header: Sequence[str], rows: Sequence[Sequence[str]]) -> bytes:
    """在内存构造 csv 上传件（utf-8-sig，Excel 导出形态）。"""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(list(header))
    for row in rows:
        writer.writerow(list(row))
    return buffer.getvalue().encode("utf-8-sig")


def positive_price(row: Mapping[str, str]) -> RowIssue | None:
    """注入的业务规则桩：参考单价必须是正数（biz_rule 段）。"""
    price = str(row.get("price", "") or "").strip()
    if not price:
        return None
    try:
        value = float(price)
    except ValueError:
        return RowIssue("参考单价", ERROR_BIZ_RULE, "「参考单价」必须是数字", "请填写金额，例如 12.50")
    if value <= 0:
        return RowIssue("参考单价", ERROR_BIZ_RULE, "「参考单价」必须大于 0", "请填写正数金额，例如 12.50")
    return None


PRODUCT_HEADER = ["标题", "正文", "SKU", "产品大类", "参考单价"]


def make_service(
    hook: RecordingHook | None = None,
    *,
    config: ImportConfig | None = None,
    existing: set[str] | None = None,
    biz_validators: Sequence[Any] = (positive_price,),
) -> tuple[ImportService, RecordingHook]:
    """组装被测服务（内存桩全链路，无 mock 截断）。"""
    hook = hook if hook is not None else RecordingHook()
    service = ImportService(
        StubSchemaProvider(),
        hook,
        config=config,
        existing_fingerprints=existing,
        biz_validators=biz_validators,
    )
    return service, hook


def valid_row(title: str = "标题A", content: str = "正文A", sku: str = "SKU-001",
              category: str = "电子", price: str = "12.50") -> list[str]:
    """一行合法 product 数据。"""
    return [title, content, sku, category, price]


async def run_job(service: ImportService, **kwargs: Any) -> Any:
    """建档并同步执行到终态（单事件循环内直接 await）。"""
    job = service.create_job(**kwargs)
    return await service.process(job.job_id)


# --------------------------------------------------------------------------- 建档校验（40001 人话）


class TestCreateJobValidation:
    """上传四重校验：格式/大小/kb_type/行数。"""

    def test_bad_format(self) -> None:
        service, _ = make_service()
        with pytest.raises(ImportValidationError, match="仅支持 xlsx / csv"):
            service.create_job(file_name="清单.txt", file_bytes=b"x", kb_type="product")

    def test_too_large(self) -> None:
        service, _ = make_service(config=ImportConfig(max_file_mb=0))  # 0MB 上限
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row()])
        with pytest.raises(ImportValidationError, match="文件大小超过"):
            service.create_job(file_name="清单.csv", file_bytes=data, kb_type="product")

    def test_unknown_kb_type(self) -> None:
        service, _ = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row()])
        with pytest.raises(ImportValidationError, match="不支持的知识库类型"):
            service.create_job(file_name="清单.csv", file_bytes=data, kb_type="oops")

    def test_bad_import_mode_and_status(self) -> None:
        service, _ = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row()])
        with pytest.raises(ImportValidationError, match="导入模式"):
            service.create_job(file_name="清单.csv", file_bytes=data, kb_type="product", import_mode="replace")
        with pytest.raises(ImportValidationError, match="入库状态"):
            service.create_job(file_name="清单.csv", file_bytes=data, kb_type="product", target_status="live")

    def test_over_max_rows(self) -> None:
        service, _ = make_service(config=ImportConfig(max_rows=2))
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row("t1"), valid_row("t2"), valid_row("t3")])
        with pytest.raises(ImportValidationError, match="不超过 2 行"):
            service.create_job(file_name="清单.csv", file_bytes=data, kb_type="product")

    def test_empty_file(self) -> None:
        service, _ = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [])
        with pytest.raises(ImportValidationError, match="没有数据行"):
            service.create_job(file_name="清单.csv", file_bytes=data, kb_type="product")

    def test_created_job_defaults(self) -> None:
        service, _ = make_service()
        job = service.create_job(
            file_name="清单.csv", file_bytes=build_csv_bytes(PRODUCT_HEADER, [valid_row()]), kb_type="product",
        )
        assert job.status == STATUS_QUEUED
        assert job.status_history == [STATUS_QUEUED]
        assert job.total_rows == 1
        assert job.import_mode == "skip" and job.target_status == "draft"


# --------------------------------------------------------------------------- 成功导入（xlsx 主路径）


class TestImportSuccess:
    """§9 成功路径：三行有效 xlsx → done、条目创建动作、导入留痕。"""

    def test_success_flow(self) -> None:
        service, hook = make_service()
        data = build_xlsx_bytes(
            PRODUCT_HEADER,
            [valid_row("音箱X1", "续航12小时", "SKU-001", "电子", "99.00"),
             valid_row("椅子Y2", "实木框架", "SKU-002", "家居", "45.50"),
             valid_row("外套Z3", "防风面料", "SKU-003", "服装", "30.00")],
        )

        async def scenario() -> Any:
            return await run_job(service, file_name="产品库.xlsx", file_bytes=data, kb_type="product")

        job = asyncio.run(scenario())
        assert job.status == STATUS_DONE
        assert (job.valid_rows, job.error_rows, job.dup_rows) == (3, 0, 0)
        assert job.processed_rows == job.total_rows == 3
        assert job.progress_pct == 100
        assert job.status_history == [STATUS_QUEUED, STATUS_RUNNING, STATUS_DONE]
        assert job.item_ids == [1, 2, 3]
        # 创建动作契约：kb_type/status/标题正文/metadata 分离/指纹携带
        assert [item.title for item in hook.items] == ["音箱X1", "椅子Y2", "外套Z3"]
        assert all(item.kb_type == "product" and item.status == "draft" for item in hook.items)
        assert all(item.overwrite is False for item in hook.items)
        assert hook.items[0].metadata == {"sku": "SKU-001", "category": "电子", "price": "99.00"}
        assert all(len(item.fingerprint) == 32 for item in hook.items)  # md5 十六进制
        # 导入留痕：处理量/成功/失败/查重计数 + excel 源格式口径
        logs = service.log_store.for_job(job.job_id)
        assert len(logs) == 1
        log = logs[0]
        assert (log.total_rows, log.valid_rows, log.error_rows, log.dup_rows) == (3, 3, 0, 0)
        assert log.source_format == "excel" and log.stage == "import"
        assert log.status == STATUS_DONE and log.item_ids == [1, 2, 3]

    def test_target_status_published_passthrough(self) -> None:
        service, hook = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row()])
        job = asyncio.run(run_job(
            service, file_name="p.csv", file_bytes=data, kb_type="product", target_status="published",
        ))
        assert job.status == STATUS_DONE
        assert hook.items[0].status == "published"

    def test_csv_leading_zero_preserved(self) -> None:
        """§9 test_csv_leading_zero：CSV 中 SKU='0012' 读入仍为字符串（前导零不丢）。"""
        service, hook = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row("带前导零", "正文", "0012")])
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        assert job.status == STATUS_DONE
        assert hook.items[0].metadata["sku"] == "0012"


# --------------------------------------------------------------------------- 部分成功与错误行报告


class TestPartialSuccess:
    """§9 test_partial_success / 全失败：坏行只报告不入库、好行照常入库。"""

    def test_partial_success_mixed_errors(self) -> None:
        service, hook = make_service()
        data = build_xlsx_bytes(
            PRODUCT_HEADER,
            [
                valid_row("好行1", "正文1", "SKU-001"),                       # 行2 有效
                ["缺SKU的行", "正文", "", "电子", "1.00"],                    # 行3 必填缺失
                ["坏选项", "正文", "SKU-002", "五金", "1.00"],               # 行4 选项越界
                ["负价格", "正文", "SKU-003", "电子", "-5"],                  # 行5 业务规则
                valid_row("好行2", "正文2", "SKU-004"),                       # 行6 有效
                valid_row("好行1", "正文1", "SKU-001"),                       # 行7 与行2 完全重复（文件内查重）
            ],
        )
        job = asyncio.run(run_job(service, file_name="p.xlsx", file_bytes=data, kb_type="product"))
        assert job.status == STATUS_DONE                    # 部分成功：任务不整单失败
        assert (job.valid_rows, job.error_rows, job.dup_rows) == (2, 4, 1)
        assert job.processed_rows == 6
        assert len(hook.items) == 2 and [i.title for i in hook.items] == ["好行1", "好行2"]

        by_row = {error.row_no: error for error in job.errors}
        # 行3：必填缺失（schema_violation，人话定位到列）
        assert by_row[3].field_label == "SKU*"
        assert by_row[3].error_code == ERROR_SCHEMA_VIOLATION
        assert "必填" in by_row[3].message and "SKU" in by_row[3].message
        assert by_row[3].suggestion
        # 行4：选项越界（含可选值建议）
        assert by_row[4].error_code == ERROR_SCHEMA_VIOLATION
        assert "电子 / 家居 / 服装" in by_row[4].suggestion
        # 行5：业务规则
        assert by_row[5].error_code == ERROR_BIZ_RULE and by_row[5].field_label == "参考单价"
        # 行7：文件内查重，行号回指首个创建行
        assert by_row[7].error_code == ERROR_DUPLICATE
        assert "与第 2 行内容重复" in by_row[7].message
        # 错误行携带原始快照（报告 sheet2 回看）
        assert by_row[3].raw_row.get("标题") == "缺SKU的行"

    def test_error_report_xlsx_content(self) -> None:
        """错误行报告：错误明细 + 原始数据标注两个 sheet，行号升序。"""
        service, _ = make_service()
        data = build_csv_bytes(
            PRODUCT_HEADER,
            [["坏行1", "正文", "", "电子", "1.00"], valid_row("好行", "正文", "SKU-009"),
             ["坏行2", "", "SKU-010", "电子", "1.00"]],
        )
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        report = build_error_report_xlsx(job)
        workbook = load_workbook(io.BytesIO(report))
        assert workbook.sheetnames == ["错误明细", "原始数据（错误已标注）"]
        detail_rows = [[cell.value for cell in row] for row in workbook["错误明细"].iter_rows()]
        assert detail_rows[0] == ["行号", "列", "错误类型", "原因", "修改建议"]
        assert [row[0] for row in detail_rows[1:]] == [2, 4]      # 行号升序
        assert detail_rows[1][1] == "SKU*"                        # 出错列中文名
        assert detail_rows[2][1] == "正文*"
        raw_rows = [[cell.value for cell in row] for row in workbook["原始数据（错误已标注）"].iter_rows()]
        assert raw_rows[0][-2:] == ["导入结果", "原因"]
        assert raw_rows[1][0] == "坏行1" and raw_rows[1][-2] == "未导入"

    def test_all_failed(self) -> None:
        """全失败：任务仍 done（0 有效行不是系统故障），报告覆盖全部行。"""
        service, hook = make_service()
        data = build_csv_bytes(
            PRODUCT_HEADER,
            [["坏1", "正文", "", "电子", "1.00"], ["坏2", "正文", "SKU", "五金", "1.00"]],
        )
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        assert job.status == STATUS_DONE
        assert (job.valid_rows, job.error_rows) == (0, 2)
        assert hook.items == []
        assert job.progress_pct == 100
        assert service.log_store.for_job(job.job_id)[0].valid_rows == 0

    def test_formula_injection_in_report(self) -> None:
        """§9 test_formula_injection：原始值 '=cmd(...)' 等在报告中前置撇号。"""
        service, _ = make_service()
        data = build_csv_bytes(
            PRODUCT_HEADER,
            [["=SUM(A1:A9)", "正文", "", "电子", "1.00"],
             ["+1+1", "@cmd|' /C calc'!A0", "SKU-77", "电子", "abc"]],
        )
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        workbook = load_workbook(io.BytesIO(build_error_report_xlsx(job)))
        raw_sheet = workbook["原始数据（错误已标注）"]
        values = [str(cell.value) for row in raw_sheet.iter_rows() for cell in row]
        assert "'=SUM(A1:A9)" in values           # 前置单引号，不作为公式
        assert "'+1+1" in values and "'@cmd|' /C calc'!A0" in values
        # 明细 sheet 的原因/建议文案同样过 sanitize
        assert all(not str(v).startswith(("=", "+", "@")) for v in values if isinstance(v, str) and v)


# --------------------------------------------------------------------------- 查重（md5 指纹联动）


class TestDedup:
    """§9 test_dup_in_file / test_dup_db_skip / overwrite：文件内 + 库内查重。"""

    def test_in_file_dup_with_md5_normalize(self) -> None:
        """文件内查重：md5 规范化口径（去首尾空白）——' X ' 与 'X' 判重。"""
        service, hook = make_service()
        data = build_csv_bytes(
            PRODUCT_HEADER,
            [valid_row(" 音箱X1 ", " 正文A ", "SKU-001"), valid_row("音箱X1", "正文A", "SKU-001")],
        )
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        assert job.dup_rows == 1
        assert any("与第 2 行内容重复" in e.message for e in job.errors)
        assert len(hook.items) == 1

    def test_dup_fingerprint_knowledge_scoped(self) -> None:
        """指纹含 kb_type：同内容跨库不判重（各库独立查重口径）。"""
        keys = ("title", "content", "sku")
        assert dedup_fingerprint("product", {"title": "T", "content": "C"}, keys) != \
            dedup_fingerprint("customer", {"title": "T", "content": "C"}, keys)

    def test_db_dup_skip_mode(self) -> None:
        """库内已存在 + skip：记 duplicate、不动旧条目（不触发创建动作）。"""
        keys = ("title", "content", "sku")
        existing = {dedup_fingerprint("product", {"title": "T1", "content": "C1", "sku": "SKU-001"}, keys)}
        service, hook = make_service(existing=existing)
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row("T1", "C1", "SKU-001")])
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        assert job.status == STATUS_DONE
        assert (job.valid_rows, job.dup_rows) == (0, 1)
        assert hook.items == []                     # skip 不动旧条目
        assert any("系统中已存在" in e.message for e in job.errors)
        assert any("覆盖导入" in e.suggestion for e in job.errors)

    def test_db_dup_overwrite_mode(self) -> None:
        """库内已存在 + overwrite：转覆盖创建动作（overwrite=True），不记错误行。"""
        keys = ("title", "content", "sku")
        existing = {dedup_fingerprint("product", {"title": "T1", "content": "C1", "sku": "SKU-001"}, keys)}
        service, hook = make_service(existing=existing)
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row("T1", "C1", "SKU-001")])
        job = asyncio.run(run_job(
            service, file_name="p.csv", file_bytes=data, kb_type="product", import_mode="overwrite",
        ))
        assert (job.valid_rows, job.error_rows, job.dup_rows) == (1, 0, 0)
        assert len(hook.items) == 1 and hook.items[0].overwrite is True

    def test_created_fingerprints_registered_to_library(self) -> None:
        """创建成功的指纹登记入库指纹集：第二个任务导入同内容判库内重复。"""
        service, _ = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row("T1", "C1", "SKU-001")])
        job = asyncio.run(run_job(service, file_name="a.csv", file_bytes=data, kb_type="product"))
        assert job.valid_rows == 1
        data2 = build_csv_bytes(PRODUCT_HEADER, [valid_row("T1", "C1", "SKU-001")])
        job2 = asyncio.run(run_job(service, file_name="b.csv", file_bytes=data2, kb_type="product"))
        assert job2.dup_rows == 1 and job2.valid_rows == 0


# --------------------------------------------------------------------------- 分片、水位续跑与降级


class TestChunkResumeAndDegrade:
    """§9 test_resume_from_watermark / test_chunk_db_error_degrade。"""

    def test_resume_from_watermark(self) -> None:
        """1002 行、500 行分片：第 600 行创建系统故障 → failed 水位 599；
        修复后续跑从第 600 行起，无重复创建，最终计数正确。"""
        hook = RecordingHook(fail_rows={"标题600"}, transient=False)
        service, hook = make_service(hook)
        header = ["标题", "正文", "公司名称"]
        rows = [[f"标题{i}", f"正文{i}", f"公司{i:04d}"] for i in range(1, 1003)]
        data = build_csv_bytes(header, rows)

        async def scenario() -> Any:
            job = service.create_job(file_name="大文件.csv", file_bytes=data, kb_type="customer")
            first = await service.process(job.job_id)      # 第 2 分片内故障中断
            assert first.status == STATUS_FAILED
            assert "导入中断" in first.error_detail
            # 水位：第 1 分片 500 行 + 第 2 分片已即时定稿的 99 行（行 600 抛错止）
            assert first.processed_rows == 599
            assert len(hook.items) == 599
            assert first.status_history == [STATUS_QUEUED, STATUS_RUNNING, STATUS_FAILED]
            hook.fail_rows.clear()                        # 「修复」存储故障
            second = await service.process(job.job_id)    # 水位续跑
            assert second is first
            return second

        job = asyncio.run(scenario())
        assert job.status == STATUS_DONE
        assert job.status_history == [STATUS_QUEUED, STATUS_RUNNING, STATUS_FAILED, STATUS_RUNNING, STATUS_DONE]
        assert job.processed_rows == job.total_rows == 1002
        assert (job.valid_rows, job.error_rows) == (1002, 0)
        assert len(hook.items) == 1002                     # 无重复创建（幂等）
        assert job.item_ids == list(range(1, 1003))
        # 留痕两条：failed 中断一条 + done 完成一条（全程可回溯）
        logs = service.log_store.for_job(job.job_id)
        assert [log.status for log in logs] == [STATUS_FAILED, STATUS_DONE]
        assert logs[1].valid_rows == 1002

    def test_chunk_transient_degrade(self) -> None:
        """瞬时故障重试耗尽 → 该行降级 internal 错误行，任务照常 done。"""
        hook = RecordingHook(fail_rows={"行2"})
        service, hook = make_service(hook, config=ImportConfig(retry_per_chunk=1))
        data = build_csv_bytes(
            PRODUCT_HEADER,
            [valid_row("行1", "正文", "SKU-001"), valid_row("行2", "正文", "SKU-002"),
             valid_row("行3", "正文", "SKU-003")],
        )
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        assert job.status == STATUS_DONE                   # 不整单失败
        assert (job.valid_rows, job.error_rows) == (2, 1)  # 行2 持续瞬时故障降级
        internal = job.errors[0]
        assert internal.error_code == ERROR_INTERNAL and internal.field_label == "整行"
        assert "系统繁忙" in internal.message and internal.suggestion
        assert [i.title for i in hook.items] == ["行1", "行3"]  # 重试轮次内其余行照常创建

    def test_chunking_boundary(self) -> None:
        """500 行分片：1200 行 → 3 个分片边界推进（水位中途可见 running）。"""
        service, _ = make_service(config=ImportConfig(chunk_rows=500))
        header = ["标题", "正文", "公司名称"]
        rows = [[f"标题{i}", f"正文{i}", f"公司{i:04d}"] for i in range(1, 1201)]
        data = build_csv_bytes(header, rows)
        snapshots: list[int] = []

        async def scenario() -> Any:
            job = service.create_job(file_name="p.csv", file_bytes=data, kb_type="customer")
            task = asyncio.get_running_loop().create_task(service.process(job.job_id))
            while not task.done():                          # 在分片让步点观察水位推进
                await asyncio.sleep(0)
                if job.status == STATUS_RUNNING and job.processed_rows not in snapshots:
                    snapshots.append(job.processed_rows)
            return await task

        job = asyncio.run(scenario())
        assert job.status == STATUS_DONE and job.valid_rows == 1200
        # 采样点恒含两个分片边界水位（500/1000），且不出现分片外的中间值
        midpoints = [value for value in snapshots if value > 0]
        assert midpoints == [500, 1000]


# --------------------------------------------------------------------------- 异步任务模型


class TestAsyncStateMachine:
    """submit（fire-and-forget）→ running → done：状态机轨迹与终态等待。"""

    def test_submit_and_wait(self) -> None:
        service, hook = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row("t1", "c1", "SKU-001"), valid_row("t2", "c2", "SKU-002")])

        async def scenario() -> Any:
            job = service.create_job(file_name="p.csv", file_bytes=data, kb_type="product")
            service.submit(job.job_id)                       # 异步启动（上传接口同路径）
            final = await service.wait_for(job.job_id)
            assert final is job
            return final

        job = asyncio.run(scenario())
        assert job.status == STATUS_DONE
        assert job.status_history == [STATUS_QUEUED, STATUS_RUNNING, STATUS_DONE]
        assert len(hook.items) == 2

    def test_process_idempotent_after_done(self) -> None:
        """done 后重复 process：直接返回，不重复创建（幂等）。"""
        service, hook = make_service()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row()])
        job = asyncio.run(run_job(service, file_name="p.csv", file_bytes=data, kb_type="product"))
        again = asyncio.run(service.process(job.job_id))
        assert again.status == STATUS_DONE and len(hook.items) == 1
        assert len(service.log_store.for_job(job.job_id)) == 1  # 留痕不重复写
