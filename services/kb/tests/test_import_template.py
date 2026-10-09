# 知识产权归甲方所有（COM-IP-01）——AI 外贸销售智能体系统交付物，许可与再分发依双方约定。
"""导入模板生成器单元测试（FR-KB-021 / SR-KB-021 §9 test_template_columns_and_dropdown）。"""
from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook

from services.kb.core.import_service.template import (
    DATA_SHEET_NAME,
    TIPS_SHEET_NAME,
    build_workbook,
    column_specs,
    display_label,
    render_csv,
    render_csv_bytes,
    sanitize_cell,
    workbook_to_bytes,
)


# --------------------------------------------------------------------------- 测试桩：注册表（鸭子类型）


@dataclass(frozen=True)
class StubField:
    """注册表字段桩：模板生成器按鸭子类型读取其属性。"""

    key: str
    label: str
    required: bool = False
    options: tuple[str, ...] = ()
    help: str = ""
    example: Any = ""
    widget: str = ""


class StubSchemaProvider:
    """Schema 注册表桩（product 库）：与 FR-KB-020 真实注册表同形。

    每实例持有独立字段表副本：单个用例内改写字段（注入恶意示例值等）
    不会污染其它用例。
    """

    FIELDS: dict[str, tuple[StubField, ...]] = {
        "product": (
            StubField("sku", "SKU", required=True, help="产品唯一货号，如 SKU-001", example="SKU-001"),
            StubField("category", "产品大类", options=("电子", "家居", "服装"), example="电子"),
            StubField("price", "参考单价", help="美元，正数，如 12.50", example="12.50"),
        ),
    }

    def __init__(self) -> None:
        self._fields: dict[str, tuple[StubField, ...]] = {
            kb: fields for kb, fields in self.FIELDS.items()
        }

    def fields(self, kb_type: str) -> tuple[StubField, ...]:
        try:
            return self._fields[kb_type]
        except KeyError:
            raise KeyError(f"未注册的知识库类型：{kb_type}") from None


# --------------------------------------------------------------------------- sanitize_cell（防公式注入）


class TestSanitizeCell:
    """§8 / §6.3：以 = + - @ 开头的字符串前置单引号转义。"""

    def test_formula_prefixes_escaped(self) -> None:
        assert sanitize_cell("=SUM(A1:A9)") == "'=SUM(A1:A9)"
        assert sanitize_cell("+1+1") == "'+1+1"
        assert sanitize_cell("-2+3") == "'-2+3"
        assert sanitize_cell("@cmd|' /C calc'") == "'@cmd|' /C calc'"

    def test_normal_values_untouched(self) -> None:
        assert sanitize_cell("普通文本") == "普通文本"
        assert sanitize_cell("") == ""
        assert sanitize_cell("1=1") == "1=1"  # 危险字符不在开头不转义
        assert sanitize_cell("加减号中文") == "加减号中文"
        assert sanitize_cell(12) == 12  # 非字符串原样返回
        assert sanitize_cell(None) is None


# --------------------------------------------------------------------------- 列规格归一化


class TestColumnSpecs:
    """列规格：基础列在前 + 注册表字段；必填星号口径。"""

    def test_base_columns_first_then_registry(self) -> None:
        specs = column_specs("product", StubSchemaProvider())
        assert [spec.key for spec in specs] == ["title", "content", "sku", "category", "price"]
        assert [display_label(spec) for spec in specs] == [
            "标题*", "正文*", "SKU*", "产品大类", "参考单价",
        ]

    def test_dynamic_options_override_static(self) -> None:
        # 动态选项（api/taxonomy 来源列）覆盖静态 options 进入下拉——与后端校验同源
        specs = column_specs("product", StubSchemaProvider(), {"category": ("美妆", "食品")})
        assert specs[3].options == ("美妆", "食品")


# --------------------------------------------------------------------------- xlsx 模板结构


class TestBuildWorkbook:
    """xlsx 模板：数据 sheet（表头/示例/下拉/冻结）+ 填写说明 sheet。"""

    def test_sheet_names_and_header(self) -> None:
        workbook = build_workbook("product", StubSchemaProvider())
        assert workbook.sheetnames == [DATA_SHEET_NAME, TIPS_SHEET_NAME]
        sheet = workbook[DATA_SHEET_NAME]
        header = [cell.value for cell in sheet[1]]
        assert header == ["标题*", "正文*", "SKU*", "产品大类", "参考单价"]

    def test_example_row_present_and_sanitized(self) -> None:
        workbook = build_workbook("product", StubSchemaProvider())
        sheet = workbook[DATA_SHEET_NAME]
        example = [cell.value for cell in sheet[2]]
        assert example[0] == "便携蓝牙音箱 X1"  # 基础列示例
        assert example[2] == "SKU-001"

    def test_example_row_formula_escaped(self) -> None:
        # 示例值以 = 开头时转义为文本（防公式注入）
        provider = StubSchemaProvider()
        provider._fields["product"] = (
            StubField("evil", "公式示例", example="=cmd|' /C calc'!A0"),
        ) + provider._fields["product"]
        workbook = build_workbook("product", provider)
        value = workbook[DATA_SHEET_NAME].cell(row=2, column=3).value
        assert value == "'=cmd|' /C calc'!A0"

    def test_dropdown_for_option_column(self) -> None:
        workbook = build_workbook("product", StubSchemaProvider())
        sheet = workbook[DATA_SHEET_NAME]
        validations = sheet.data_validations.dataValidation
        assert len(validations) == 1  # 仅「产品大类」选项列有下拉
        validation = validations[0]
        assert validation.type == "list"
        assert validation.formula1 == '"电子,家居,服装"'
        assert list(validation.sqref.ranges)[0].coord == "D2:D10001"  # 第 4 列、2~10001 行
        assert validation.errorTitle == "填写不符合要求"
        assert "产品大类" in validation.error

    def test_no_dropdown_for_free_text(self) -> None:
        # 自由文本列（标题/正文/SKU/单价）不生成下拉
        workbook = build_workbook("product", StubSchemaProvider())
        sheet = workbook[DATA_SHEET_NAME]
        formulas = [dv.formula1 for dv in sheet.data_validations.dataValidation]
        assert all("SKU-001" not in f for f in formulas)

    def test_tips_sheet_rows(self) -> None:
        workbook = build_workbook("product", StubSchemaProvider())
        tips = workbook[TIPS_SHEET_NAME]
        rows = [[cell.value for cell in row] for row in tips.iter_rows()]
        assert rows[0] == ["列名", "必填", "格式/说明", "示例"]
        assert len(rows) == 6  # 表头 + 5 列（2 基础 + 3 注册表）
        sku_row = rows[3]
        assert sku_row == ["SKU*", "必填", "产品唯一货号，如 SKU-001", "SKU-001"]
        category_row = rows[4]
        assert category_row[1] == "选填"

    def test_workbook_roundtrip_openable_read_only(self) -> None:
        # 生成的模板可被导入解析端以 read-only 模式读回（生成/解析闭环）
        data = workbook_to_bytes(build_workbook("product", StubSchemaProvider()))
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        try:
            rows = list(workbook.worksheets[0].iter_rows(values_only=True))
        finally:
            workbook.close()
        assert rows[0] == ("标题*", "正文*", "SKU*", "产品大类", "参考单价")


# --------------------------------------------------------------------------- csv 模板


class TestRenderCsv:
    """csv 模板：仅列头 + 示例行，utf-8-sig 带 BOM。"""

    def test_csv_header_and_example(self) -> None:
        text = render_csv("product", StubSchemaProvider())
        lines = text.splitlines()
        assert lines[0].split(",") == ["标题*", "正文*", "SKU*", "产品大类", "参考单价"]
        assert "SKU-001" in lines[1]

    def test_csv_bytes_with_bom(self) -> None:
        data = render_csv_bytes("product", StubSchemaProvider())
        assert data[:3] == b"\xef\xbb\xbf"  # Excel 双击直开不乱码
        assert "SKU-001" in data.decode("utf-8-sig")

    def test_csv_formula_example_escaped(self) -> None:
        provider = StubSchemaProvider()
        provider._fields["product"] = (
            StubField("evil", "公式示例", example="+1+1"),
        ) + provider._fields["product"]
        text = render_csv("product", provider)
        assert "'+1+1" in text.splitlines()[1]
