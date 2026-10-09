# 知识产权归甲方所有（COM-IP-01）——AI 外贸销售智能体系统交付物。
# 许可与再分发须依据双方合同约定执行（COM-IP-02 乙方零使用权）。
"""数据清洗报告聚合器（FR-KB-024，SR-KB-024 全章）。

知识流水线（解析→清洗→分块→向量化→入库）的**只读旁路报表环节**：聚合
ingest_pipeline_log 留痕生成报告，不侵入、不阻塞清洗主链路（决策 D1）。

五要素口径（决策 D2，与 SR-KB-011 §4 action_detail 键一一对应）：
- 处理量   = parse 段最新留痕的 blocks_in 之和（块）+ 处理文件数（parse 成功文件数）；
- 去重量   = dedup 段最新留痕的 dropped_exact + dropped_near 之和；
- 纠错量   = correct 段最新留痕的 corrections 之和；
- 脱敏量   = mask 段最新留痕的 masked_blocks 之和（兼容 0.1 迭代清洗 stats 的
             masked_findings 键，两键并存时以 masked_blocks 为准）；
- 遗留问题 = 最新留痕仍为 failed/pending 的问题记录（含历史窗口遗留），
             三分类：解析失败 / 清洗冲突 / 人工待审，各带处置建议（§6.2）。

取数口径（§6.1 latest）：每 (source_file, stage) 只取最新一条留痕——
FR-KB-017 幂等重放成功后旧失败行即失效（test_replay_supersede）。

留痕行结构沿用迭代 0.1 的 ingest_pipeline_log 语义（stage / action_detail / output）：
action_detail 约定键 = status(success|failed|pending) / batch_id / kb_type /
blocks_in / dropped_exact / dropped_near / corrections / masked_blocks / needs_review /
error_code / error_msg；output 为该阶段产物归档键（如 p0/clean-output/{md5}.json）。

导出：xlsx 三 sheet（汇总 / 文件明细 / 遗留问题清单，§6.3）+ md 同口径镜像；
触发：批次完成回调 + 每日定时钩子（§6.4），同批次/同窗口幂等只生成一份。

依据文档：docs/07-实现方案/KB/SR-KB-024-数据清洗报告.md；
         docs/03-技术文档/02-开发规范.md §3（中文注释 / 全量类型注解）。
"""
from __future__ import annotations

import re
import threading
import uuid
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Protocol

# ── openpyxl 守护导入：聚合/分类/触发不依赖它，仅 xlsx 渲染需要（核心依赖，
#    环境未安装时给出带中文提示的明确错误；对应测试用 importorskip 自动跳过）──
try:  # pragma: no cover - 分支取决于环境，两侧均非永久路径
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    _OPENPYXL_AVAILABLE = True
except ImportError:  # pragma: no cover
    Workbook = None  # type: ignore[assignment]
    PatternFill = None  # type: ignore[assignment]
    Font = None  # type: ignore[assignment]
    _OPENPYXL_AVAILABLE = False

# ── 常量：阶段 / 状态枚举（对齐 ingest_pipeline_log.stage 的 CHECK 枚举）──
STAGE_PARSE = "parse"
STAGE_DEDUP = "dedup"
STAGE_CORRECT = "correct"
STAGE_MASK = "mask"
ALL_STAGES: tuple[str, ...] = ("parse", "dedup", "correct", "mask", "chunk", "embed", "index")

STATUS_SUCCESS = "success"
STATUS_FAILED = "failed"
STATUS_PENDING = "pending"
BAD_STATUSES: tuple[str, ...] = (STATUS_FAILED, STATUS_PENDING)

# 终态中文化（文件明细 sheet「终态」列，FR-KB-023 中文化口径）
TERMINAL_CN: dict[str, str] = {STATUS_SUCCESS: "成功", STATUS_FAILED: "失败", STATUS_PENDING: "待审"}

# 五要素中文名（汇总 sheet 行序即此序）
ELEMENT_PROCESSED = "处理量"
ELEMENT_FILES = "处理文件数"
ELEMENT_DEDUP = "去重量"
ELEMENT_CORRECT = "纠错量"
ELEMENT_MASK = "脱敏量"
FIVE_ELEMENTS: tuple[str, ...] = (
    ELEMENT_PROCESSED,
    ELEMENT_DEDUP,
    ELEMENT_CORRECT,
    ELEMENT_MASK,
)
"""五要素（SRS §5.4 原文序：处理量、去重量、纠错量、脱敏量、遗留问题清单）。

处理文件数与遗留问题为口径补充列（遗留问题单列于问题清单 sheet 及汇总行）。"""

# PII 复扫正则（NFR-SEC-01 硬门：渲染文本零命中才允许写入报告，与清洗引擎默认规则同式）
_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+")
_PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d{1,4}[\s\-]?)?\d{7,13}(?!\d)")
_IDNUM_RE = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")


class PiiLeakError(ValueError):
    """报告文本中检出疑似明文 PII（NFR-SEC-01 硬门拦截）。"""


def scan_plaintext_pii(text: str) -> str | None:
    """扫描文本，命中明文 PII 时返回命中的串（用于错误信息），否则返回 None。"""
    for pattern in (_EMAIL_RE, _PHONE_RE, _IDNUM_RE):
        hit = pattern.search(text)
        if hit is not None:
            return hit.group(0)
    return None


def _assert_no_pii(text: str, *, where: str) -> None:
    """渲染前 PII 硬门：问题摘要/建议等动态文本零命中才放行（SR-KB-024 §8）。"""
    hit = scan_plaintext_pii(text)
    if hit is not None:  # pragma: no cover - 防御分支，正常数据不携带 PII
        raise PiiLeakError(f"{where} 检出疑似明文 PII（{hit}），报告拒绝渲染（NFR-SEC-01）")


# --------------------------------------------------------------------------- 留痕数据源（SR §4 / §6.1）


@dataclass(frozen=True)
class IngestLogRecord:
    """ingest_pipeline_log 一行的内存形态（迭代 0.1 留痕契约：stage/action_detail/output）。

    id 为留痕自增序号（latest 取数与重放失效判定均以 id 最大者为最新）；
    action_detail 约定键见模块 docstring；output 为阶段产物归档键（可空）。
    """

    source_file: str
    source_format: str
    stage: str
    action_detail: Mapping[str, Any] = field(default_factory=dict)
    id: int = 0  # 0 表示由存储分配（对齐 PG IDENTITY 语义）
    created_at: datetime | None = None  # None 表示由存储按追加时刻生成
    output: str = ""

    @property
    def status(self) -> str:
        """本条留痕状态（约定键 status；缺失视为 success——0.1 迭代成功留痕可省略）。"""
        return str(self.action_detail.get("status", STATUS_SUCCESS))

    @property
    def batch_id(self) -> str:
        """批次号（FR-KB-021 导入批次；历史缺失按 '-' 归并）。"""
        return str(self.action_detail.get("batch_id", "-"))

    @property
    def kb_type(self) -> str:
        """知识库类型（历史缺失按 'unmapped' 归并，可由 dir_mapper 回填）。"""
        return str(self.action_detail.get("kb_type", "unmapped"))

    @property
    def created(self) -> datetime:
        """留痕时间（缺失时按当前 UTC 补全，聚合侧不应依赖该补全）。"""
        return self.created_at or datetime.now(timezone.utc)


class IngestLogStore(Protocol):
    """清洗留痕存储协议（生产为 PG ingest_pipeline_log 只读查询；报告侧只读，决策 D1）。"""

    def iter_logs(
        self,
        *,
        batch_id: str | None = None,
        window_from: datetime | None = None,
        window_to: datetime | None = None,
    ) -> Iterable[IngestLogRecord]:
        """按批次号（action_detail.batch_id）或时间窗 [window_from, window_to) 迭代留痕。"""
        ...  # pragma: no cover


class InMemoryIngestLogStore:
    """内存留痕存储（单元测试 / 单机演示用；线程安全，追加分配自增 id）。"""

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._rows: list[IngestLogRecord] = []
        self._max_id = 0

    @property
    def rows(self) -> list[IngestLogRecord]:
        """内部行列表活引用：仅供测试只读检视，生产代码勿用。"""
        return self._rows

    def append(self, record: IngestLogRecord) -> int:
        """追加一条留痕（id=0 时分配自增 id；created_at 缺省取当前 UTC），返回 id。"""
        with self._lock:
            self._max_id += 1
            rid = record.id if record.id > 0 else self._max_id
            self._max_id = max(self._max_id, rid)
            created = record.created_at or datetime.now(timezone.utc)
            self._rows.append(
                IngestLogRecord(
                    source_file=record.source_file,
                    source_format=record.source_format,
                    stage=record.stage,
                    action_detail=dict(record.action_detail),
                    id=rid,
                    created_at=created,
                    output=record.output,
                )
            )
            return rid

    def extend(self, records: Iterable[IngestLogRecord]) -> None:
        """批量追加（测试便捷方法）。"""
        for r in records:
            self.append(r)

    def iter_logs(
        self,
        *,
        batch_id: str | None = None,
        window_from: datetime | None = None,
        window_to: datetime | None = None,
    ) -> Iterable[IngestLogRecord]:
        """按批次 / 时间窗过滤（批次优先；与 ReportScope 语义一致）。"""
        with self._lock:
            snapshot = list(self._rows)
        for r in snapshot:
            if batch_id is not None and r.batch_id != batch_id:
                continue
            if window_from is not None and (r.created_at or datetime.now(timezone.utc)) < window_from:
                continue
            if window_to is not None and (r.created_at or datetime.now(timezone.utc)) >= window_to:
                continue
            yield r


def _as_int(detail: Mapping[str, Any], key: str) -> int:
    """计数键安全取整（缺失/None/非法值 → 0；兼容 0.1 迭代字符串计数字段）。"""
    raw = detail.get(key)
    if raw is None:
        return 0
    try:
        return int(raw)
    except (TypeError, ValueError):  # pragma: no cover - 防御分支：脏数据按 0 计
        return 0


def _masked_of(detail: Mapping[str, Any]) -> int:
    """脱敏计数：masked_blocks 优先，历史键 masked_findings 兜底（0.1 迭代清洗 stats）。"""
    if "masked_blocks" in detail:
        return _as_int(detail, "masked_blocks")
    return _as_int(detail, "masked_findings")


# --------------------------------------------------------------------------- 报告范围与聚合（SR §6.1）


@dataclass(frozen=True)
class ReportScope:
    """报告范围：batch_id（批次）与时间窗二选一（接口层已互斥校验，§5 手动生成入参）。"""

    batch_id: str | None = None
    window_from: datetime | None = None
    window_to: datetime | None = None

    def __post_init__(self) -> None:
        has_batch = self.batch_id is not None and self.batch_id != ""
        has_window = self.window_from is not None or self.window_to is not None
        if has_batch and has_window:
            raise ValueError("报告范围互斥：batch_id 与时间窗不得同时给出（E_RPT_SCOPE_CONFLICT）")
        if not has_batch and not has_window:
            raise ValueError("报告范围必填：batch_id 或时间窗二选一")
        if has_window and (self.window_from is None or self.window_to is None):
            raise ValueError("时间窗必须同时给出 window_from 与 window_to")
        if has_window and self.window_from >= self.window_to:  # type: ignore[operator]
            raise ValueError("时间窗非法：window_from 必须早于 window_to")

    @property
    def is_batch(self) -> bool:
        """是否批次范围。"""
        return self.batch_id is not None and self.batch_id != ""

    def key(self) -> str:
        """幂等键（同批次/同窗口至多一份报告，§4 唯一约束的内存等价物）。"""
        if self.is_batch:
            return f"batch:{self.batch_id}"
        assert self.window_from is not None and self.window_to is not None
        return f"win:{self.window_from:%Y%m%d%H%M%S}-{self.window_to:%Y%m%d%H%M%S}"

    def label(self) -> str:
        """报告展示用范围标签（批次号或时间窗区间）。"""
        if self.is_batch:
            return f"批次 {self.batch_id}"
        assert self.window_from is not None and self.window_to is not None
        return f"{self.window_from:%Y-%m-%d %H:%M} ~ {self.window_to:%Y-%m-%d %H:%M}"


@dataclass
class FileDetail:
    """文件明细行（明细 sheet 逐文件计数，§16.1 数据完整性抽样核对入口）。

    可变聚合容器：聚合器逐 stage 累加字段、收敛终态；产出后即视为只读
    （与 audit.AuditRow 同款「内存实现可变」约定）。
    """

    source_file: str
    source_format: str
    kb_type: str
    blocks_in: int = 0
    dedup: int = 0
    corrections: int = 0
    masked: int = 0
    terminal: str = STATUS_SUCCESS  # 终态：任一阶段最新留痕非 success 即非成功（failed 优先于 pending）

    @property
    def terminal_cn(self) -> str:
        """终态中文化。"""
        return TERMINAL_CN.get(self.terminal, self.terminal)

    def row(self) -> tuple[Any, ...]:
        """明细 sheet 行（列序与渲染表头一致）。"""
        return (
            self.source_file,
            self.source_format,
            self.kb_type,
            self.blocks_in,
            self.dedup,
            self.corrections,
            self.masked,
            self.terminal_cn,
        )


@dataclass(frozen=True)
class AggSnapshot:
    """聚合快照（kb_clean_report.stats JSONB 的内存形态，决策 D3 报告快照化）。"""

    scope: ReportScope
    totals: dict[str, int]  # 五要素 + 处理文件数
    by_kb: dict[str, dict[str, int]]  # 分库矩阵（仅含出现过的库，报告瘦身）
    by_format: dict[str, dict[str, int]]  # 分格式矩阵（仅含出现过的格式）
    file_details: tuple[FileDetail, ...]
    ruleset_version: str  # 清洗规则集版本（快照可复现的前提）

    def element(self, name: str) -> int:
        """按五要素中文名取数（缺失键按 0，防御直接构造的快照）。"""
        return int(self.totals.get(name, 0))


def _latest_by_file_stage(records: Sequence[IngestLogRecord]) -> dict[tuple[str, str], IngestLogRecord]:
    """每 (source_file, stage) 取 id 最大的一条（§6.1 latest 口径：重放成功后旧失败行失效）。"""
    latest: dict[tuple[str, str], IngestLogRecord] = {}
    for r in records:
        k = (r.source_file, r.stage)
        cur = latest.get(k)
        if cur is None or r.id > cur.id:
            latest[k] = r
    return latest


def _element_row() -> dict[str, int]:
    """要素行的零值骨架（矩阵用）。"""
    return {
        ELEMENT_PROCESSED: 0,
        ELEMENT_DEDUP: 0,
        ELEMENT_CORRECT: 0,
        ELEMENT_MASK: 0,
    }


class CleanReportAggregator:
    """报告聚合器：按批次/时间窗只读聚合留痕，输出五要素 + 分库×分格式矩阵（§6.1）。"""

    def __init__(self, store: IngestLogStore) -> None:
        self.store = store

    def _fetch(self, scope: ReportScope) -> list[IngestLogRecord]:
        """按范围取留痕（批次优先，不受跨天影响）。"""
        if scope.is_batch:
            return list(self.store.iter_logs(batch_id=scope.batch_id))
        assert scope.window_from is not None and scope.window_to is not None
        return list(
            self.store.iter_logs(window_from=scope.window_from, window_to=scope.window_to)
        )

    def aggregate(self, scope: ReportScope) -> AggSnapshot:
        """聚合五要素与矩阵（要素只取其归属 stage 的最新留痕，决策 D2）。"""
        latest = _latest_by_file_stage(self._fetch(scope))

        totals: dict[str, int] = {
            ELEMENT_PROCESSED: 0,
            ELEMENT_FILES: 0,
            ELEMENT_DEDUP: 0,
            ELEMENT_CORRECT: 0,
            ELEMENT_MASK: 0,
        }
        by_kb: dict[str, dict[str, int]] = {}
        by_format: dict[str, dict[str, int]] = {}
        details: dict[str, FileDetail] = {}
        ruleset_versions: set[str] = set()

        def _kb_row(kb: str) -> dict[str, int]:
            return by_kb.setdefault(kb, _element_row())

        def _fmt_row(fmt: str) -> dict[str, int]:
            return by_format.setdefault(fmt, _element_row())

        def _detail(rec: IngestLogRecord) -> FileDetail:
            return details.setdefault(
                rec.source_file,
                FileDetail(
                    source_file=rec.source_file,
                    source_format=rec.source_format,
                    kb_type=rec.kb_type,
                    terminal=STATUS_SUCCESS,
                ),
            )

        for (source_file, stage), rec in latest.items():
            detail_rec = rec.action_detail
            d = _detail(rec)
            if rec.status == STATUS_FAILED or (
                rec.status == STATUS_PENDING and d.terminal == STATUS_SUCCESS
            ):
                # 终态取最差：failed 优先于 pending，success 不覆盖
                d.terminal = rec.status
            if "mask_ruleset_version" in detail_rec:
                ruleset_versions.add(str(detail_rec["mask_ruleset_version"]))
            if "ruleset_version" in detail_rec:
                ruleset_versions.add(str(detail_rec["ruleset_version"]))

            # 五要素归属 stage 计数（决策 D2：要素只取其归属 stage 的行）
            if stage == STAGE_PARSE:
                blocks = _as_int(detail_rec, "blocks_in")
                totals[ELEMENT_PROCESSED] += blocks
                if rec.status == STATUS_SUCCESS:
                    totals[ELEMENT_FILES] += 1
                _kb_row(rec.kb_type)[ELEMENT_PROCESSED] += blocks
                _fmt_row(rec.source_format)[ELEMENT_PROCESSED] += blocks
                d.blocks_in = blocks
            elif stage == STAGE_DEDUP:
                dropped = _as_int(detail_rec, "dropped_exact") + _as_int(detail_rec, "dropped_near")
                totals[ELEMENT_DEDUP] += dropped
                _kb_row(rec.kb_type)[ELEMENT_DEDUP] += dropped
                _fmt_row(rec.source_format)[ELEMENT_DEDUP] += dropped
                d.dedup = dropped
            elif stage == STAGE_CORRECT:
                corrections = _as_int(detail_rec, "corrections")
                totals[ELEMENT_CORRECT] += corrections
                _kb_row(rec.kb_type)[ELEMENT_CORRECT] += corrections
                _fmt_row(rec.source_format)[ELEMENT_CORRECT] += corrections
                d.corrections = corrections
            elif stage == STAGE_MASK:
                masked = _masked_of(detail_rec)
                totals[ELEMENT_MASK] += masked
                _kb_row(rec.kb_type)[ELEMENT_MASK] += masked
                _fmt_row(rec.source_format)[ELEMENT_MASK] += masked
                d.masked = masked

        return AggSnapshot(
            scope=scope,
            totals=totals,
            by_kb=by_kb,
            by_format=by_format,
            file_details=tuple(
                details[k] for k in sorted(details, key=lambda s: (len(s), s))
            ),
            ruleset_version="|".join(sorted(ruleset_versions)) or "unknown",
        )


# --------------------------------------------------------------------------- 遗留问题三分类（SR §6.2）


CLASS_PARSE_FAILED = "parse_failed"
CLASS_CLEAN_CONFLICT = "clean_conflict"
CLASS_MANUAL_REVIEW = "manual_review"

CLASS_CN: dict[str, str] = {
    CLASS_PARSE_FAILED: "解析失败",
    CLASS_CLEAN_CONFLICT: "清洗冲突",
    CLASS_MANUAL_REVIEW: "人工待审",
}
"""三分类中文名（问题清单 sheet「分类」列，FR-KB-023 中文化口径）。"""

ISSUE_RULES: tuple[tuple[tuple[str, ...], tuple[str, ...], str, str], ...] = (
    (
        ("*",),
        (STATUS_PENDING,),
        CLASS_MANUAL_REVIEW,
        "人工待审：脱敏边界/近似去重复核样本，请在管理后台复核后回写结论",
    ),
    (
        (STAGE_PARSE,),
        (STATUS_FAILED,),
        CLASS_PARSE_FAILED,
        "解析失败：检查原件可读性，调整 OCR/ASR 参数重试；不可解析转人工登记",
    ),
    (
        (STAGE_DEDUP, STAGE_CORRECT, STAGE_MASK, "chunk", "embed", "index"),
        (STATUS_FAILED,),
        CLASS_CLEAN_CONFLICT,
        "清洗/流水线冲突：按 error_code 修正规则集或依赖后，经 FR-KB-017 幂等重放本文件",
    ),
)
"""分类规则（顺序即优先级）：stage × status → 分类 + 处置建议（文案随规则版本化入 app_config）。"""


@dataclass(frozen=True)
class CleanIssue:
    """遗留问题登记行（kb_clean_issue 的内存形态；log_id 为幂等锚点——一条留痕至多一个问题）。"""

    log_id: int
    batch_id: str
    source_file: str
    source_format: str
    kb_type: str
    stage: str
    issue_class: str
    description: str  # 仅计数/错误码/规则版本，不含 PII（NFR-SEC-01）
    suggestion: str  # 处置建议（ISSUE_RULES 文案）
    status: str = "open"  # open | resolved | wont_fix
    resolved_note: str = ""

    @property
    def class_cn(self) -> str:
        """分类中文名。"""
        return CLASS_CN.get(self.issue_class, self.issue_class)

    def row(self) -> tuple[Any, ...]:
        """问题清单 sheet 行（列序与渲染表头一致）。"""
        return (
            self.log_id,
            self.class_cn,
            self.source_file,
            self.stage,
            self.description,
            self.suggestion,
            self.status,
            self.batch_id,
        )


class IssueBuilder:
    """遗留问题构建：扫描 failed/pending 最新留痕 → 三分类 → 登记问题表（§6.2）。

    - 幂等：以 log_id 为锚点（一条留痕至多一个问题），重复同步不重复登记；
    - 自动闭合：留痕已翻转为 success 的旧问题 → resolved（note=重放成功自动闭合），
      遗留清单随报告滚动更新（决策 D5）。
    """

    def classify(self, stage: str, status: str) -> tuple[str, str]:
        """stage × status → (分类, 处置建议)；规则顺序即优先级，未知失败兜底清洗冲突。"""
        for stages, statuses, cls, tip in ISSUE_RULES:
            if status in statuses and ("*" in stages or stage in stages):
                return cls, tip
        return CLASS_CLEAN_CONFLICT, ISSUE_RULES[2][3]  # 兜底：未知失败按清洗冲突

    def _describe(self, rec: IngestLogRecord) -> str:
        """问题描述（≤200 字符；仅错误码/摘要/规则版本，不含 PII）。"""
        if rec.status == STATUS_PENDING:
            base = f"[pending] {rec.action_detail.get('error_msg', '人工待审挂起')}"
        else:
            code = rec.action_detail.get("error_code", "unknown")
            msg = rec.action_detail.get("error_msg", "")
            base = f"[{code}] {msg}" if msg else f"[{code}] 阶段执行失败"
        return base[:200]

    def sync_issues(
        self,
        store: IngestLogStore,
        scope: ReportScope,
        registry: dict[int, CleanIssue],
    ) -> int:
        """同步问题登记（登记新增 + 自动闭合），返回全量 open 计数（含历史窗口遗留）。"""
        aggregator_latest = _latest_by_file_stage(
            list(CleanReportAggregator(store)._fetch(scope))
        )
        # 1) 登记最新仍为 failed/pending 的留痕（log_id 幂等锚点）
        for (source_file, stage), rec in aggregator_latest.items():
            if rec.status not in BAD_STATUSES:
                continue
            if rec.id in registry and registry[rec.id].status == "open":
                continue  # 已登记且仍 open：幂等跳过
            cls, tip = self.classify(stage, rec.status)
            _assert_no_pii(f"{source_file} {self._describe(rec)}", where="问题描述")
            registry[rec.id] = CleanIssue(
                log_id=rec.id,
                batch_id=rec.batch_id if rec.batch_id != "-" else (scope.batch_id or "-"),
                source_file=source_file,
                source_format=rec.source_format,
                kb_type=rec.kb_type,
                stage=stage,
                issue_class=cls,
                description=self._describe(rec),
                suggestion=tip,
                status="open",
            )
        # 2) 自动闭合：登记项对应 (文件, 阶段) 的最新留痕已翻转为 success → resolved
        latest_index = aggregator_latest
        for issue in list(registry.values()):
            if issue.status != "open":
                continue
            current = None
            for (source_file, stage), rec in latest_index.items():
                if source_file == issue.source_file and stage == issue.stage:
                    current = rec
                    break
            if current is not None and current.status not in BAD_STATUSES:
                closed = CleanIssue(
                    log_id=issue.log_id,
                    batch_id=issue.batch_id,
                    source_file=issue.source_file,
                    source_format=issue.source_format,
                    kb_type=issue.kb_type,
                    stage=issue.stage,
                    issue_class=issue.issue_class,
                    description=issue.description,
                    suggestion=issue.suggestion,
                    status="resolved",
                    resolved_note="重放成功自动闭合",
                )
                registry[issue.log_id] = closed
        return sum(1 for i in registry.values() if i.status == "open")


# --------------------------------------------------------------------------- 报告对象与渲染（SR §6.3）


REPORT_KIND_BATCH = "batch"
REPORT_KIND_DAILY = "daily"
REPORT_KIND_MANUAL = "manual"

_SUMMARY_SHEET = "汇总"
_DETAIL_SHEET = "文件明细"
_ISSUE_SHEET = "遗留问题清单"
XLSX_SHEET_NAMES: tuple[str, ...] = (_SUMMARY_SHEET, _DETAIL_SHEET, _ISSUE_SHEET)
"""xlsx 三 sheet 名（§6.3：汇总 / 文件明细 / 遗留问题清单）。"""

_DETAIL_HEADERS: tuple[str, ...] = (
    "来源文件",
    "格式",
    "知识库",
    "处理量",
    "去重量",
    "纠错量",
    "脱敏量",
    "终态",
)
_ISSUE_HEADERS: tuple[str, ...] = (
    "编号",
    "分类",
    "来源文件",
    "阶段",
    "问题描述",
    "处置建议",
    "状态",
    "首次发现批次",
)

MAX_DETAIL_ROWS = 5000
"""明细 sheet 行上限（§7 report.clean.max_detail_rows 默认值；超出截断并在汇总 sheet 标注）。"""


@dataclass
class CleaningReport:
    """一次报告生成结果（kb_clean_report 行的内存形态 + 聚合快照，决策 D3）。"""

    report_uid: str
    report_kind: str  # batch | daily | manual
    scope: ReportScope
    snapshot: AggSnapshot
    issues: tuple[CleanIssue, ...]  # 全量 open（含历史窗口遗留，决策 D5）
    issue_new: int  # 本期新增遗留
    issue_carry: int  # 历史遗留（open 全集 − 本期新增）
    created_at: datetime
    created_by: str = "system"
    detail_truncated: bool = False  # 明细超上限截断标记（汇总 sheet 标注）

    @property
    def scope_label(self) -> str:
        """范围标签（批次或时间窗）。"""
        return self.scope.label()

    @property
    def open_issue_count(self) -> int:
        """遗留问题总数。"""
        return len(self.issues)

    def issue_by_class(self) -> dict[str, int]:
        """遗留问题按三分类计数（中文键）。"""
        out: dict[str, int] = {cn: 0 for cn in CLASS_CN.values()}
        for i in self.issues:
            out[i.class_cn] = out.get(i.class_cn, 0) + 1
        return out


class ReportRenderer:
    """中文报告渲染：xlsx 为 P0 交付物主格式（三 sheet）；md 用于评审邮件，
    数值与 xlsx 同源同口径（§6.3）。"""

    def __init__(self, workdir: str | Path | None = None) -> None:
        self.workdir = Path(workdir) if workdir is not None else Path("out") / "clean-reports"
        self.workdir.mkdir(parents=True, exist_ok=True)

    # ── xlsx ──────────────────────────────────────────────────────────────

    def render_xlsx(self, report: CleaningReport) -> Path:
        """渲染三 sheet xlsx，返回文件路径（本地渲染完成后再归档 OSS，无半成品文件）。"""
        if not _OPENPYXL_AVAILABLE:  # pragma: no cover - 环境分支
            raise ImportError(
                "openpyxl 未安装：数据清洗报告 xlsx 导出依赖 openpyxl（核心依赖），"
                "请安装后重试；聚合/分类/md 导出不受影响"
            )
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = _SUMMARY_SHEET
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(name="微软雅黑", bold=True, color="FFFFFF")

        def _style_header(cell: Any) -> None:
            cell.fill = header_fill
            cell.font = header_font

        def _write_table(
            sheet: Any, start_row: int, headers: Sequence[str], rows: Sequence[Sequence[Any]]
        ) -> int:
            """写一张带样式表头的表，返回下一可用行号。"""
            for col, h in enumerate(headers, start=1):
                cell = sheet.cell(row=start_row, column=col, value=h)
                _style_header(cell)
            for r_off, row in enumerate(rows, start=1):
                for col, v in enumerate(row, start=1):
                    sheet.cell(row=start_row + r_off, column=col, value=v)
            return start_row + len(rows) + 2

        # 汇总 sheet：报告元信息 → 五要素 → 分库矩阵 → 分格式矩阵
        r = 1
        for label, value in (
            ("报告类型", report.report_kind),
            ("批次/时间窗", report.scope_label),
            ("生成时间", report.created_at.isoformat(timespec="seconds")),
            ("清洗规则集版本", report.snapshot.ruleset_version),
        ):
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=value)
            r += 1
        r += 1
        snap = report.snapshot
        element_rows: list[tuple[Any, ...]] = [
            (k, v) for k, v in snap.totals.items()
        ] + [("遗留问题（本期新增/历史遗留）", f"{report.issue_new} / {report.issue_carry}")]
        r = _write_table(ws, r, ("指标", "数值"), element_rows)
        r = _write_table(
            ws, r, ("知识库", ELEMENT_PROCESSED, ELEMENT_DEDUP, ELEMENT_CORRECT, ELEMENT_MASK),
            [(kb, *row.values()) for kb, row in snap.by_kb.items()],
        )
        r = _write_table(
            ws, r, ("来源格式", ELEMENT_PROCESSED, ELEMENT_DEDUP, ELEMENT_CORRECT, ELEMENT_MASK),
            [(fmt, *row.values()) for fmt, row in snap.by_format.items()],
        )
        if report.detail_truncated:
            ws.cell(
                row=r, column=1,
                value=f"注：文件明细超过 {MAX_DETAIL_ROWS} 行，已截断（完整口径见 stats 快照）",
            )

        # 文件明细 sheet（上限截断）
        details = report.snapshot.file_details[:MAX_DETAIL_ROWS]
        ws_detail = wb.create_sheet(_DETAIL_SHEET)
        _write_table(ws_detail, 1, _DETAIL_HEADERS, [d.row() for d in details])

        # 遗留问题清单 sheet（全量 open + 处置建议，分类中文）
        ws_issue = wb.create_sheet(_ISSUE_SHEET)
        _write_table(ws_issue, 1, _ISSUE_HEADERS, [i.row() for i in report.issues])

        path = self.workdir / f"{report.report_uid}.xlsx"
        wb.save(path)
        return path

    # ── md ────────────────────────────────────────────────────────────────

    def render_md(self, report: CleaningReport) -> str:
        """渲染 Markdown 报告（与 xlsx 同源同口径：五要素/矩阵/遗留问题逐一对应）。"""
        snap = report.snapshot
        lines: list[str] = [
            "# 数据清洗报告",
            "",
            f"- 报告类型：{report.report_kind}",
            f"- 批次/时间窗：{report.scope_label}",
            f"- 生成时间：{report.created_at.isoformat(timespec='seconds')}",
            f"- 清洗规则集版本：{snap.ruleset_version}",
            "",
            "## 五要素汇总",
            "",
            "| 指标 | 数值 |",
            "|---|---|",
        ]
        for k, v in snap.totals.items():
            lines.append(f"| {k} | {v} |")
        lines.append(f"| 遗留问题（本期新增/历史遗留） | {report.issue_new} / {report.issue_carry} |")
        lines += ["", "## 分知识库", "", "| 知识库 | 处理量 | 去重量 | 纠错量 | 脱敏量 |", "|---|---|---|---|---|"]
        for kb, row in snap.by_kb.items():
            cells = (kb, row[ELEMENT_PROCESSED], row[ELEMENT_DEDUP], row[ELEMENT_CORRECT], row[ELEMENT_MASK])
            lines.append("| " + " | ".join(str(c) for c in cells) + " |")
        lines += ["", "## 分来源格式", "", "| 来源格式 | 处理量 | 去重量 | 纠错量 | 脱敏量 |", "|---|---|---|---|---|"]
        for fmt, row in snap.by_format.items():
            cells = (fmt, row[ELEMENT_PROCESSED], row[ELEMENT_DEDUP], row[ELEMENT_CORRECT], row[ELEMENT_MASK])
            lines.append("| " + " | ".join(str(c) for c in cells) + " |")
        lines += ["", "## 遗留问题清单", ""]
        if not report.issues:
            lines.append("（本期无遗留问题）")
        else:
            lines.append("| 分类 | 来源文件 | 阶段 | 问题描述 | 处置建议 | 状态 |")
            lines.append("|---|---|---|---|---|---|")
            for i in report.issues:
                desc = i.description.replace("|", "\\|")
                tip = i.suggestion.replace("|", "\\|")
                lines.append(f"| {i.class_cn} | {i.source_file} | {i.stage} | {desc} | {tip} | {i.status} |")
        if report.detail_truncated:
            lines += ["", f"> 注：文件明细超过 {MAX_DETAIL_ROWS} 行，已截断。"]
        return "\n".join(lines) + "\n"


# --------------------------------------------------------------------------- 触发编排（SR §6.4）


class CleanReportService:
    """报告编排：聚合 → 问题同步 → 渲染 → 登记，batch/daily/manual 三类触发（§6.4）。

    幂等：同 (kind, scope) 至多一份报告——重复触发直接返回既有报告
    （生产为 Redis 锁 + 表级唯一约束双重幂等，此为进程内等价物）；
    只读旁路：本服务故障不影响知识入库主流程（决策 D1）。
    """

    def __init__(
        self,
        store: IngestLogStore,
        *,
        renderer: ReportRenderer | None = None,
        issue_registry: dict[int, CleanIssue] | None = None,
        aggregator: CleanReportAggregator | None = None,
        issue_builder: IssueBuilder | None = None,
        now_provider: Callable[[], datetime] | None = None,
    ) -> None:
        self.store = store
        self.renderer = renderer or ReportRenderer()
        self.issue_registry: dict[int, CleanIssue] = (
            issue_registry if issue_registry is not None else {}
        )
        self.aggregator = aggregator or CleanReportAggregator(store)
        self.issue_builder = issue_builder or IssueBuilder()
        self._now_provider = now_provider or (lambda: datetime.now(timezone.utc))
        self._reports: dict[str, CleaningReport] = {}  # (kind, scope.key()) → 报告
        self.generate_count = 0  # 实际执行生成次数（幂等测试观测用）

    # ── 生成主流程 ────────────────────────────────────────────────────────

    def generate(self, kind: str, scope: ReportScope, actor: str = "system") -> CleaningReport:
        """生成（或幂等返回）一份报告：聚合 → 问题同步 → 截断标记 → 登记。"""
        cache_key = f"{kind}:{scope.key()}"
        existing = self._reports.get(cache_key)
        if existing is not None:
            return existing  # 幂等：同批次/同窗口重复触发返回既有报告
        known_before = set(self.issue_registry)  # 本期新增 = 本次同步期间新登记者
        snap = self.aggregator.aggregate(scope)
        open_cnt = self.issue_builder.sync_issues(self.store, scope, self.issue_registry)
        new_ids = {
            rid
            for rid in set(self.issue_registry) - known_before
            if self.issue_registry[rid].status == "open" and self._in_scope(self.issue_registry[rid], scope)
        }
        report = CleaningReport(
            report_uid=f"r-{uuid.uuid4().hex[:12]}",
            report_kind=kind,
            scope=scope,
            snapshot=snap,
            issues=tuple(
                i for i in self.issue_registry.values() if i.status == "open"
            ),
            issue_new=len(new_ids),
            issue_carry=open_cnt - len(new_ids),
            created_at=self._now_provider(),
            created_by=actor,
            detail_truncated=len(snap.file_details) > MAX_DETAIL_ROWS,
        )
        self._reports[cache_key] = report
        self.generate_count += 1
        return report

    def _in_scope(self, issue: CleanIssue, scope: ReportScope) -> bool:
        """问题是否属于本期范围（本期新增判定：批次号或时间窗内命中）。"""
        if scope.is_batch:
            return issue.batch_id == scope.batch_id
        assert scope.window_from is not None and scope.window_to is not None
        for rec in self.store.iter_logs():
            if rec.id == issue.log_id:
                created = rec.created_at or self._now_provider()
                return scope.window_from <= created < scope.window_to
        return False

    # ── 触发器一：批次完成回调（§6.4 scan_batches 的回调形态）──────────────

    def on_batch_complete(self, batch_id: str, actor: str = "system") -> CleaningReport:
        """批次完成回调：批次闭合（全部终态且静默）由调用方判定后触发本钩子（幂等）。"""
        return self.generate(REPORT_KIND_BATCH, ReportScope(batch_id=batch_id), actor=actor)

    # ── 触发器二：每日定时钩子（§6.4 APScheduler cron 0 2 * * * 的入口）────

    def on_daily_timer(
        self, day: date | None = None, *, actor: str = "system"
    ) -> CleaningReport:
        """每日定时钩子：默认聚合昨日全天窗口 [昨日 00:00, 今日 00:00)，同窗幂等。

        生产形态为 APScheduler cron（report.clean.daily_cron，默认 0 2 * * *，
        时区 Asia/Shanghai）；本钩子为其唯一入口，重复触发返回既有报告。
        """
        target = day if day is not None else self._now_provider().date() - timedelta(days=1)
        window_from = datetime.combine(target, time.min, tzinfo=timezone.utc)
        window_to = datetime.combine(target + timedelta(days=1), time.min, tzinfo=timezone.utc)
        return self.generate(
            REPORT_KIND_DAILY,
            ReportScope(window_from=window_from, window_to=window_to),
            actor=actor,
        )

    # ── 手动生成（§5 POST /generate 的服务入口）──────────────────────────

    def generate_manual(
        self, *, batch_id: str | None = None, actor: str = "system",
        window_from: datetime | None = None, window_to: datetime | None = None,
    ) -> CleaningReport:
        """手动生成（batch_id 与时间窗二选一；互斥校验在 ReportScope 构造期完成）。"""
        scope = ReportScope(
            batch_id=batch_id, window_from=window_from, window_to=window_to
        )
        return self.generate(REPORT_KIND_MANUAL, scope, actor=actor)

    # ── 查询 ──────────────────────────────────────────────────────────────

    def list_reports(self) -> list[CleaningReport]:
        """已生成报告列表（生成序）。"""
        return list(self._reports.values())

    def export_files(self, report: CleaningReport) -> tuple[Path, str]:
        """导出 xlsx + md 双格式（生产另上传 OSS p0/clean-reports/ 归档，决策 D3）。"""
        return self.renderer.render_xlsx(report), self.renderer.render_md(report)
