# 知识产权归甲方所有（COM-IP-01）——AI 外贸销售智能体系统交付物，许可与再分发依双方合同约定。
"""§16.1 KPI 预跑样例语料生成器（迭代 0.3）：构建 tools/precheck/sample_corpus/ 下的小型样例素材。

产出八个素材文件，覆盖六大知识库与三种可离线解析格式（docx / xlsx / pdf）：

| 文件 | kb_type | 格式 | 用途 |
|---|---|---|---|
| product_产品手册_ST3200.docx | product | word | 产品库样例（FR-KB-001） |
| product_参数表.xlsx | product | excel | 产品参数表样例（表格块链路） |
| product_产品目录.pdf | product | pdf | PDF 解析链路样例 |
| customer_分级标准.docx | customer | word | 客户库样例（FR-KB-002） |
| skill_议价话术.docx | skill | word | 技能库样例（FR-KB-003） |
| process_报价SOP.docx | process | word | 流程库样例（FR-KB-004） |
| negative_丢单复盘.docx | negative_case | word | 负面案例库样例（FR-KB-005） |
| training_新人第一课.docx | training | word | 带教素材库样例（FR-KB-006） |

内容约定（评测集命中前提，FakeEmbeddingClient 稀疏路按 unicode \\w 分词）：
- 中文锚点句以标点/空格切分出独立词串，评测 query 复用同一词串保证词元命中；
- 不含明文 PII 与七位以上数字串（避免默认脱敏规则改写锚点）；
- 单位直接使用标准单位（千克/厘米），避开纠错归一化改写。

注意：样例仅供 §16.1 KPI 预跑（开发自验）使用；生产验收必须替换为甲方真实
素材与不少于一百条的人工评测集（见 tools/precheck/phase0_kpi.py 报告口径）。

用法：python3 tools/precheck/sample_corpus/build_samples.py [输出目录]
"""
from __future__ import annotations

import sys
from pathlib import Path

#: 默认输出目录（脚本所在目录）
DEFAULT_OUT = Path(__file__).resolve().parent

# ── 各素材正文（(标题, [段落]) 形态；标题进 Heading 1，段落进 Normal）────────

DOCX_SAMPLES: dict[str, tuple[str, list[str]]] = {
    "product_产品手册_ST3200.docx": (
        "ST-3200 电视支架产品手册",
        [
            "ST-3200 为壁挂式电视支架，最大承载85千克。",
            "机身采用冷轧钢板，表面静电喷涂工艺，适配三十七至七十英寸电视。",
            "认证情况：CE 认证与 SGS 认证均已通过。",
            "安装孔距支持 200x200 至 600x400 毫米范围，随附安装模板与膨胀螺栓。",
        ],
    ),
    "customer_分级标准.docx": (
        "客户分级标准说明",
        [
            "A 级客户为月采购超过五千件的核心经销商，享受优先排产权益。",
            "B 级客户为月采购五百至五千件的稳定客户，按标准交期安排生产。",
            "C 级客户为偶发询盘客户，首单建议样品确认后再批量。",
            "D 级客户为长期无互动或存在拖账风险的客户，新订单须预付定金。",
        ],
    ),
    "skill_议价话术.docx": (
        "砍价应对话术要点",
        [
            "客户要求降价百分之十时，先确认目标数量再回应阶梯价格。",
            "让步必须换取对等条件：订单量、付款方式或交期弹性三选一。",
            "底线坚守话术：这个价格已经是我们对经销商的授权底线。",
            "赠品策略：价格不动，改赠安装配件或延长质保至五年。",
        ],
    ),
    "process_报价SOP.docx": (
        "P03 报价节点操作规程",
        [
            "收到询盘后两个工作日内完成正式报价单的制作与发出。",
            "报价单必须包含型号、规格、FOB 单价、最小起订量与有效期。",
            "价格数据一律从产品库结构化字段取数，禁止口头随意承诺。",
            "超出授权折扣范围时，先提交销售主管审批再回复客户。",
        ],
    ),
    "negative_丢单复盘.docx": (
        "丢单复盘：交期承诺失误",
        [
            "案例概况：向欧洲经销商承诺十五天交货，实际排产需要二十八天。",
            "直接后果：客户错过本地销售旺季，转向竞品供应商下单。",
            "根因分析：报价阶段未与生产部门确认排产能力，交期口径凭经验估计。",
            "改进措施：交期承诺必须走排产确认流程，加急需求单独评估。",
        ],
    ),
    "training_新人第一课.docx": (
        "新人第一课：产品体系入门",
        [
            "课程目标：能说出公司三大产品线与各自的主打型号。",
            "支架产品线主打型号为 ST-3200，面向客厅电视安装场景。",
            "云台产品线主打型号为 GP-100，面向摄影稳定拍摄场景。",
            "课后练习：写出三个型号的卖点各一条，明天晨会分享。",
        ],
    ),
}

#: Excel 样例：sheet 名 → 二维表（首行为表头）
XLSX_SAMPLES: dict[str, dict[str, list[list[str]]]] = {
    "product_参数表.xlsx": {
        "参数表": [
            ["参数", "数值", "说明"],
            ["承载重量", "85千克", "最大负载"],
            ["适配尺寸", "三十七至七十英寸", "电视对角线"],
            ["安装孔距", "600x400", "毫米，最大支持"],
            ["本体重量", "4.6千克", "含包装为七千克"],
        ],
    },
}

#: PDF 样例：文件名 → (标题行, [正文行])
PDF_SAMPLES: dict[str, tuple[str, list[str]]] = {
    "product_产品目录.pdf": (
        "ST-3200 支架产品目录",
        [
            "ST-3200 壁挂支架，最大承载85千克，通过 CE 与 SGS 双认证。",
            "支持三十七至七十英寸电视，安装孔距最大 600x400 毫米。",
            "包装规格：每箱两台，毛重十四千克，出口唛头按订单印刷。",
        ],
    ),
}


def build_docx(path: Path, title: str, paragraphs: list[str]) -> None:
    """生成 Word 样例：Heading 1 标题 + Normal 段落（word_parser 按样式识别标题）。"""
    from docx import Document

    document = Document()
    document.add_heading(title, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def build_xlsx(path: Path, sheets: dict[str, list[list[str]]]) -> None:
    """生成 Excel 样例：每 sheet 一张二维表（excel_parser 按逻辑表整块发射）。"""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, rows in sheets.items():
        sheet = workbook.create_sheet(title=sheet_name)
        for row in rows:
            sheet.append(row)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(path))


def build_pdf(path: Path, title: str, lines: list[str]) -> None:
    """生成 PDF 样例：内置简体中文字形（china-s）逐行写入单页文本框。"""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    body = title + "\n" + "\n".join(lines)
    page.insert_textbox(fitz.Rect(50, 60, 545, 780), body, fontname="china-s", fontsize=12)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()


def build_all(out_dir: Path) -> list[Path]:
    """生成全部样例素材，返回产物路径清单（存在性由调用方核对）。"""
    created: list[Path] = []
    for name, (title, paragraphs) in DOCX_SAMPLES.items():
        build_docx(out_dir / name, title, paragraphs)
        created.append(out_dir / name)
    for name, sheets in XLSX_SAMPLES.items():
        build_xlsx(out_dir / name, sheets)
        created.append(out_dir / name)
    for name, (title, lines) in PDF_SAMPLES.items():
        build_pdf(out_dir / name, title, lines)
        created.append(out_dir / name)
    return created


def main(argv: list[str]) -> int:
    """命令行入口：生成样例并打印清单（重复执行幂等覆盖）。"""
    out_dir = Path(argv[1]).resolve() if len(argv) > 1 else DEFAULT_OUT
    for path in build_all(out_dir):
        print(f"已生成: {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
