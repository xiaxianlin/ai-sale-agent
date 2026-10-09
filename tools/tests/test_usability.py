# 知识产权归甲方所有（COM-IP-01）——AI 外贸销售智能体系统交付物，许可与再分发依双方合同约定。
# test_usability.py —— 后台易用性工具元测试（FR-KB-023 / NFR-USA-01 / NFR-USA-02，迭代 0.3）：
# ops_registry 解析校验 / coverage_check 缺口检出 / task_sheet 任务单结构 / zh_scan 白名单与命中。
"""工具自身的可信性测试（金样本回归）：误放行与误阻断同等是缺陷；
样例一律写入 pytest tmp_path，仓库真实状态另设门禁级断言（coverage 全支撑）。"""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from tools.usability import coverage_check, ops_registry, task_sheet, zh_scan

REPO_ROOT = Path(__file__).resolve().parents[2]
OPS_YAML = REPO_ROOT / "tools" / "usability" / "operations.yaml"

# ─────────────────────── ops_registry：受限 YAML 解析与校验 ───────────────────────


def test_miniyaml解析布尔整数引号与注释() -> None:
    """受限 YAML 子集：布尔/整数/引号剥壳/整行注释/多字段条目全部正确解析。"""
    text = (
        "# 头注释\n"
        "schema_version: 1\n"
        "operations:\n"
        "  - code: DEMO_A\n"
        "    name: 演示操作\n"
        "    api: GET /api/v1/demo\n"
        "    steps: 2\n"
        "    dangerous: false\n"
        "  - code: DEMO_B\n"
        "    name: '带引号名称'\n"
        "    api: POST /api/v1/demo/{uid}\n"
        "    steps: 3\n"
        "    dangerous: high\n"
    )
    parsed = ops_registry.parse_simple_yaml(text)
    assert parsed["schema_version"] == 1
    ops = parsed["operations"]
    assert ops[0] == {"code": "DEMO_A", "name": "演示操作", "api": "GET /api/v1/demo",
                      "steps": 2, "dangerous": False}
    assert ops[1]["dangerous"] == "high" and ops[1]["name"] == "带引号名称"


def test_load_operations真实清单结构() -> None:
    """仓库 operations.yaml：17 项、步数全部 ≤3、两处高危（删除/回滚）、api 形态合法。"""
    operations = ops_registry.load_operations(OPS_YAML)
    assert len(operations) == 17
    assert all(op.steps <= ops_registry.STEPS_LIMIT for op in operations)
    danger_codes = {op.code for op in operations if op.dangerous == "high"}
    assert danger_codes == {"KB_ITEM_DELETE", "KB_VERSION_ROLLBACK"}
    by_code = {op.code: op for op in operations}
    assert by_code["KB_SEARCH_API"].method == "POST"
    assert by_code["KB_SEARCH_API"].path == "/api/v1/knowledge/search"
    assert by_code["KB_VERSION_ROLLBACK"].danger_label == "高危（输入名称确认）"


def test_load_operations缺文件与非法登记报错(tmp_path: Path) -> None:
    """文件缺失 FileNotFoundError；操作码/api/steps 非法与重复登记 ValueError。"""
    with pytest.raises(FileNotFoundError):
        ops_registry.load_operations(tmp_path / "nowhere.yaml")

    bad_code = tmp_path / "bad_code.yaml"
    bad_code.write_text(
        "operations:\n  - code: bad-code\n    name: x\n    api: GET /a\n", encoding="utf-8"
    )
    with pytest.raises(ValueError, match="操作码"):
        ops_registry.load_operations(bad_code)

    bad_api = tmp_path / "bad_api.yaml"
    bad_api.write_text(
        "operations:\n  - code: OK_A\n    name: x\n    api: get /a\n", encoding="utf-8"
    )
    with pytest.raises(ValueError, match="api 字段不合法"):
        ops_registry.load_operations(bad_api)

    dup = tmp_path / "dup.yaml"
    dup.write_text(
        "operations:\n"
        "  - code: OK_A\n    name: x\n    api: GET /a\n"
        "  - code: OK_A\n    name: y\n    api: GET /b\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="重复登记"):
        ops_registry.load_operations(dup)

    missing_field = tmp_path / "missing.yaml"
    missing_field.write_text("operations:\n  - code: OK_A\n    api: GET /a\n", encoding="utf-8")
    with pytest.raises(ValueError, match="缺字段"):
        ops_registry.load_operations(missing_field)


# ─────────────────────── coverage_check：AST 提取与缺口检出 ───────────────────────

_SAMPLE_API = '''\
# 样例路由（测试构造，非仓库代码）
from fastapi import APIRouter

router = APIRouter(prefix="/api/v1")

_PATH = "/const-path"


@router.get("/things")
async def list_things():
    return {}


@router.post("/things/{uid}")
async def save_thing(uid: str):
    return {}


@router.post(_PATH)
async def const_path():
    return {}


def mount_extra(app, path="/internal/extra"):
    @app.get(path)
    def extra():
        return {}
'''


def _write_sample_api(tmp_path: Path) -> Path:
    api_dir = tmp_path / "api"
    api_dir.mkdir()
    (api_dir / "sample.py").write_text(_SAMPLE_API, encoding="utf-8")
    return api_dir


def test_AST提取_常量路径与工厂内端点(tmp_path: Path) -> None:
    """模块级路径常量、工厂函数内 @app.get(path)（参数默认值）都能提取。"""
    endpoints = coverage_check.extract_endpoints(_write_sample_api(tmp_path), tmp_path)
    apis = {(ep.method, ep.path) for ep in endpoints}
    assert ("GET", "/api/v1/things") in apis
    assert ("POST", "/api/v1/things/{uid}") in apis
    assert ("POST", "/api/v1/const-path") in apis
    assert ("GET", "/internal/extra") in apis  # 参数默认值解析


def test_缺口检出_缺API漏洞与超3步(tmp_path: Path) -> None:
    """声明无支撑 → [缺API]；路由未登记 → [漏洞]；steps=4 → [超3步]；main 退出码 1。"""
    api_dir = _write_sample_api(tmp_path)
    ops_yaml = tmp_path / "ops.yaml"
    ops_yaml.write_text(
        "operations:\n"
        "  - code: OK_LIST\n    name: 列表\n    api: GET /api/v1/things\n    steps: 2\n"
        "  - code: MISS_DEL\n    name: 删除\n    api: DELETE /api/v1/things/{uid}\n    steps: 2\n"
        "  - code: SLOW_OP\n    name: 超步\n    api: GET /api/v1/slow\n    steps: 4\n",
        encoding="utf-8",
    )
    result = coverage_check.check_coverage(
        ops_registry.load_operations(ops_yaml),
        coverage_check.extract_endpoints(api_dir, tmp_path),
    )
    assert not result.passed
    assert any("[缺API] MISS_DEL" in line for line in result.gaps)
    assert any("[超3步] SLOW_OP" in line for line in result.oversteps)
    assert any("[漏洞]" in line and "/api/v1/things/{uid}" in line for line in result.wild_routes)
    assert coverage_check.main(["--api-root", str(api_dir), "--ops", str(ops_yaml)]) == 1


def test_全支撑通过_与模板参数归一匹配(tmp_path: Path) -> None:
    """声明与路由一一对应 → 通过；声明 {id} 与路由 {uid} 经归一化仍视为同一端点。"""
    api_dir = _write_sample_api(tmp_path)
    ops_yaml = tmp_path / "ops.yaml"
    ops_yaml.write_text(
        "operations:\n"
        "  - code: OK_LIST\n    name: 列表\n    api: GET /api/v1/things\n    steps: 2\n"
        "  - code: OK_SAVE\n    name: 保存\n    api: POST /api/v1/things/{id}\n    steps: 2\n"
        "  - code: OK_CONST\n    name: 常量\n    api: POST /api/v1/const-path\n    steps: 1\n",
        encoding="utf-8",
    )
    result = coverage_check.check_coverage(
        ops_registry.load_operations(ops_yaml),
        coverage_check.extract_endpoints(api_dir, tmp_path),
    )
    assert result.passed, result.problems
    assert not result.wild_routes  # /internal/extra 豁免不计漏洞
    assert any("模板参数归一匹配" in row["status"] for row in result.rows)
    assert coverage_check.main(["--api-root", str(api_dir), "--ops", str(ops_yaml)]) == 0


def test_真实仓库核对_全部门禁通过() -> None:
    """门禁级断言：当前 services/kb/api 全部路由被 operations.yaml 覆盖（新增路由必须同步登记）。"""
    operations = ops_registry.load_operations(OPS_YAML)
    endpoints = coverage_check.extract_endpoints(REPO_ROOT / "services" / "kb" / "api", REPO_ROOT)
    result = coverage_check.check_coverage(operations, endpoints)
    assert result.passed, result.problems
    assert len(result.rows) == 17
    assert all("已支撑" in row["status"] for row in result.rows)


# ─────────────────────── task_sheet：任务单结构与一致性 ───────────────────────


@pytest.mark.filterwarnings("ignore:Workbook contains no default style")
def test_任务单生成_xlsx与md结构(tmp_path: Path) -> None:
    """生成两件套：20 任务行、记录表三态列/耗时/困惑点齐备、完成状态下拉约束就位。"""
    xlsx_path, md_path = task_sheet.generate(ops_registry.load_operations(OPS_YAML), tmp_path)
    assert xlsx_path.is_file() and md_path.is_file()

    workbook = load_workbook(xlsx_path)
    assert workbook.sheetnames == ["实测任务单", "实测记录表"]
    sheet = workbook["实测任务单"]
    assert [cell.value for cell in sheet[1]] == list(task_sheet.TASK_HEADERS)
    task_rows = [row for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    assert len(task_rows) == len(task_sheet.TASKS) == 20
    assert all(int(row[7]) > 0 for row in task_rows)  # 预计分钟列为正整数
    assert all(row[3] and "/api/" in row[3] for row in task_rows)  # API 支撑列非空

    record = workbook["实测记录表"]
    headers = [cell.value for cell in record[1]]
    for column in ("完成状态（完成/求助完成/失败）", "耗时（分钟）", "困惑点原话（脱敏）"):
        assert column in headers
    validations = record.data_validations.dataValidation
    assert validations and "求助完成" in validations[0].formula1

    markdown = md_path.read_text(encoding="utf-8")
    assert "实测记录表" in markdown and "困惑点原话" in markdown
    assert markdown.count("\n| T") >= 40  # 20 任务 ×（任务表 + 记录表）两处出现


def test_任务展开_操作全覆盖与高危提示() -> None:
    """20 任务引用的操作码全部合法，且 operations.yaml 每个操作至少被一个任务覆盖。"""
    operations = ops_registry.load_operations(OPS_YAML)
    rows = task_sheet.build_task_rows(operations)
    covered = {code for row in rows for code in row["ops"].split(" + ")}
    assert covered == {op.code for op in operations}
    delete_row = next(row for row in rows if row["no"] == "T09")
    assert "高危" in delete_row["danger"] and "二次确认" in delete_row["preconditions"]


def test_任务引用未登记操作码拒绝生成() -> None:
    """任务引用操作码不在清单 → ValueError（两份事实源不允许漂移）。"""
    with pytest.raises(ValueError, match="未登记操作码"):
        task_sheet.build_task_rows([])  # 空清单：T01 即引用 KB_ITEM_CREATE


def test_操作未被任务覆盖拒绝生成() -> None:
    """任务引用的操作全登记，但清单存在无任务覆盖的操作 → ValueError（任务集完整性）。"""
    operations = [*ops_registry.load_operations(OPS_YAML)]
    operations.append(
        ops_registry.Operation(
            code="KB_FUTURE_OP", name="未来迭代操作",
            api="GET /api/v1/knowledge/future", steps=1, dangerous="false",
        )
    )
    with pytest.raises(ValueError, match="未被任何任务覆盖"):
        task_sheet.build_task_rows(operations)


# ─────────────────────── zh_scan：白名单与命中 ───────────────────────


_SCAN_SAMPLE = '''\
# 样例模块（测试构造）
from fastapi import APIRouter

from services.kb.api.routes import ApiError

router = APIRouter(prefix="/api/v1")


@router.post("/orders", summary="Submit order")
async def create_order():
    raise ApiError(40001, 400, "bad input")
'''


def _zh_hit_texts(findings: list[zh_scan.Finding]) -> list[str]:
    return [finding.text for finding in findings]


def test_英文残留判定与白名单() -> None:
    """纯英文命中；含汉字放行；白名单专有名词剔除后无英文 → 放行；ok 协议放行。"""
    assert zh_scan.is_english_residue("Submit order")
    assert zh_scan.is_english_residue("please retry later")
    assert zh_scan.is_english_residue("queued | running | done")
    assert not zh_scan.is_english_residue("请输入 SKU 编号")          # 含汉字
    assert not zh_scan.is_english_residue("TopK")                     # 白名单
    assert not zh_scan.is_english_residue("UUID 与 trace_id 说明")     # 含汉字
    assert not zh_scan.is_english_residue("ok")                       # 协议成功消息


def test_扫描命中_kwarg位置参数与f字符串(tmp_path: Path) -> None:
    """summary 关键字、ApiError 第 3 位置参数、f-string 常量段均被捕获并定位行号。"""
    src = tmp_path / "src"
    src.mkdir()
    (src / "api.py").write_text(_SCAN_SAMPLE + '''

async def status(job_id: str):
    raise ApiError(40401, 404, f"job not found: {job_id}")


def ok_path():
    return _envelope("t", 0, "ok", None)
''', encoding="utf-8")
    findings = zh_scan.scan_directory([src], tmp_path)
    texts = _zh_hit_texts(findings)
    assert "Submit order" in texts
    assert "bad input" in texts
    assert any("job not found" in text for text in texts)
    assert "ok" not in texts  # 协议成功消息放行
    summary_finding = next(f for f in findings if f.text == "Submit order")
    assert summary_finding.file == "src/api.py"
    assert summary_finding.line == 9 and summary_finding.carrier == "summary"


def test_中文文案零命中与strict退出码(tmp_path: Path) -> None:
    """全中文样例零命中；--strict 命中退出码 1、无命中退出码 0。"""
    clean = tmp_path / "clean"
    clean.mkdir()
    (clean / "zh.py").write_text(
        'router.post("/x", summary="新建知识条目")\n'
        'Field(description="请填写「产品名称」")\n',
        encoding="utf-8",
    )
    assert zh_scan.scan_directory([clean], tmp_path) == []
    assert zh_scan.main(["--roots", str(clean), "--strict"]) == 0

    dirty = tmp_path / "dirty"
    dirty.mkdir()
    (dirty / "en.py").write_text('summary_value = dict(summary="Loading")\n', encoding="utf-8")
    assert zh_scan.main(["--roots", str(dirty)]) == 0            # 缺省报告形态不阻断
    assert zh_scan.main(["--roots", str(dirty), "--strict"]) == 1
