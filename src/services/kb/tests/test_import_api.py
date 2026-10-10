# 知识产权归甲方所有（COM-IP-01）——AI 外贸销售智能体系统交付物，许可与再分发依双方约定。
"""批量导入管理端 API 集成测试（FR-KB-021 / SR-KB-021 §5、§9）。

TestClient 全链路（真实 FastAPI 应用 + 内存桩服务，无 mock 截断）：
上传 202 → 轮询 done → 错误行报告下载；模板下载（xlsx/csv）；
错误码信封（40001 上传参数 / 40401 任务不存在）。

注：TestClient 上下文管理器内所有请求共用同一事件循环（portal），POST 中
submit 的后台任务在轮询间隙持续推进——与生产「上传即返回、后台导入」同构。
"""
from __future__ import annotations

import io
import time

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from services.kb.api.admin.importer import router as import_router
from services.kb.core.import_service import ImportService
from services.kb.tests.test_import_pipeline import (
    PRODUCT_HEADER,
    RecordingHook,
    build_csv_bytes,
    build_xlsx_bytes,
    make_service,
    valid_row,
)

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
IMPORT_PATH = "/api/v1/knowledge/import"
#: 轮询超时（秒）：内存桩全链路毫秒级完成，5s 仅作防抖上限
POLL_TIMEOUT_S = 5.0


def build_app(hook: RecordingHook | None = None) -> tuple[FastAPI, ImportService, RecordingHook]:
    """组装被测应用：admin 导入路由 + app.state.import_service 注入。"""
    service, hook = make_service(hook)
    app = FastAPI(title="kb-admin 导入路由测试")
    app.include_router(import_router)
    app.state.import_service = service
    return app, service, hook


def wait_done(client: TestClient, job_id: str) -> dict:
    """轮询任务直至终态（done/failed），返回 data 视图。"""
    deadline = time.monotonic() + POLL_TIMEOUT_S
    while time.monotonic() < deadline:
        body = client.get(f"{IMPORT_PATH}/{job_id}").json()
        assert body["code"] == 0
        if body["data"]["status"] in ("done", "failed"):
            return body["data"]
        time.sleep(0.02)
    pytest.fail(f"导入任务 {job_id} 在 {POLL_TIMEOUT_S}s 内未到达终态")


def upload(
    client: TestClient,
    *,
    file_name: str,
    data: bytes,
    form: dict[str, str] | None = None,
    headers: dict[str, str] | None = None,
) -> object:
    """multipart 上传辅助。"""
    return client.post(
        IMPORT_PATH,
        files={"file": (file_name, data, "application/octet-stream")},
        data=form or {"kb_type": "product"},
        headers=headers,
    )


class TestImportApiHappyPath:
    """主链路：202 受理 → 轮询 done → 报告下载。"""

    def test_upload_poll_and_download_report(self) -> None:
        app, service, hook = build_app()
        data = build_xlsx_bytes(
            PRODUCT_HEADER,
            [
                valid_row("音箱X1", "续航12小时", "SKU-001", "电子", "99.00"),
                ["缺SKU行", "正文", "", "电子", "1.00"],          # 错误行（必填缺失）
                valid_row("椅子Y2", "实木框架", "SKU-002", "家居", "45.50"),
            ],
        )
        with TestClient(app) as client:
            response = upload(client, file_name="产品库.xlsx", data=data)
            # 202 受理：统一信封 + job_id
            assert response.status_code == 202
            body = response.json()
            assert body["code"] == 0 and body["trace_id"]
            job_id = body["data"]["job_id"]
            assert body["data"]["status"] in ("queued", "running", "done")

            final = wait_done(client, job_id)
            assert final["status"] == "done"
            assert (final["valid_rows"], final["error_rows"]) == (2, 1)
            assert final["progress_pct"] == 100
            assert final["total_rows"] == 3 and final["processed_rows"] == 3
            assert final["kb_type"] == "product"
            assert final["error_report_url"] == f"{IMPORT_PATH}/{job_id}/errors"

            # 错误行报告下载：xlsx 文件流 + 中文文件名（RFC 5987）
            report = client.get(f"{IMPORT_PATH}/{job_id}/errors")
            assert report.status_code == 200
            assert report.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
            disposition = report.headers["content-disposition"]
            assert "attachment" in disposition and "filename*=UTF-8''" in disposition
            workbook = load_workbook(io.BytesIO(report.content))
            detail_rows = [[cell.value for cell in row] for row in workbook["错误明细"].iter_rows()]
            assert detail_rows[0] == ["行号", "列", "错误类型", "原因", "修改建议"]
            assert detail_rows[1][0] == 3 and detail_rows[1][1] == "SKU*"  # 出错行号/列定位

            # 条目创建动作确已发生（2 条）
            assert [item.title for item in hook.items] == ["音箱X1", "椅子Y2"]
            assert service.log_store.for_job(job_id)[0].valid_rows == 2

    def test_trace_id_passthrough(self) -> None:
        app, _, _ = build_app()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row()])
        with TestClient(app) as client:
            response = upload(
                client, file_name="p.csv", data=data, headers={"X-Trace-Id": "trace-import-001"}
            )
            assert response.json()["trace_id"] == "trace-import-001"

    def test_csv_upload_supported(self) -> None:
        app, _, hook = build_app()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row("csv行", "正文", "0012")])
        with TestClient(app) as client:
            response = upload(client, file_name="清单.csv", data=data)
            job_id = response.json()["data"]["job_id"]
            final = wait_done(client, job_id)
            assert final["status"] == "done" and final["valid_rows"] == 1
            assert hook.items[0].metadata["sku"] == "0012"     # csv 前导零保留


class TestImportApiErrors:
    """错误码信封：40001 上传参数 / 40401 任务不存在。"""

    def test_unknown_job_404(self) -> None:
        app, _, _ = build_app()
        with TestClient(app) as client:
            response = client.get(f"{IMPORT_PATH}/no-such-job")
            assert response.status_code == 404
            body = response.json()
            assert body["code"] == 40401 and body["data"] is None
            assert "不存在" in body["message"]

    def test_unknown_job_errors_404(self) -> None:
        app, _, _ = build_app()
        with TestClient(app) as client:
            response = client.get(f"{IMPORT_PATH}/no-such-job/errors")
            assert response.status_code == 404 and response.json()["code"] == 40401

    def test_bad_file_format_400(self) -> None:
        app, _, _ = build_app()
        with TestClient(app) as client:
            response = upload(client, file_name="清单.txt", data="标题\n".encode("utf-8"))
            assert response.status_code == 400
            body = response.json()
            assert body["code"] == 40001 and "仅支持 xlsx / csv" in body["message"]

    def test_bad_kb_type_400(self) -> None:
        app, _, _ = build_app()
        data = build_csv_bytes(PRODUCT_HEADER, [valid_row()])
        with TestClient(app) as client:
            response = upload(client, file_name="p.csv", data=data, form={"kb_type": "oops"})
            assert response.status_code == 400
            assert response.json()["code"] == 40001
            assert "不支持的知识库类型" in response.json()["message"]

    def test_empty_file_400(self) -> None:
        app, _, _ = build_app()
        data = build_csv_bytes(PRODUCT_HEADER, [])
        with TestClient(app) as client:
            response = upload(client, file_name="p.csv", data=data)
            assert response.status_code == 400 and "没有数据行" in response.json()["message"]


class TestImportTemplateDownload:
    """模板下载端点：xlsx（下拉+说明）/ csv（BOM）与参数校验。"""

    def test_download_xlsx_template(self) -> None:
        app, _, _ = build_app()
        with TestClient(app) as client:
            response = client.get(f"{IMPORT_PATH}/template", params={"kb_type": "product"})
            assert response.status_code == 200
            assert response.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
            workbook = load_workbook(io.BytesIO(response.content))
            assert workbook.sheetnames == ["数据", "填写说明"]
            header = [cell.value for cell in workbook["数据"][1]]
            assert header == ["标题*", "正文*", "SKU*", "产品大类", "参考单价"]

    def test_download_csv_template(self) -> None:
        app, _, _ = build_app()
        with TestClient(app) as client:
            response = client.get(f"{IMPORT_PATH}/template", params={"kb_type": "product", "format": "csv"})
            assert response.status_code == 200
            assert "csv" in response.headers["content-type"]
            assert response.content[:3] == b"\xef\xbb\xbf"       # utf-8-sig BOM
            assert "SKU*" in response.content.decode("utf-8-sig")

    def test_template_bad_params_400(self) -> None:
        app, _, _ = build_app()
        with TestClient(app) as client:
            bad_type = client.get(f"{IMPORT_PATH}/template", params={"kb_type": "oops"})
            assert bad_type.status_code == 400 and bad_type.json()["code"] == 40001
            bad_format = client.get(f"{IMPORT_PATH}/template", params={"kb_type": "product", "format": "pdf"})
            assert bad_format.status_code == 400


class TestImportApiFailurePath:
    """任务级失败经状态端点可观测（error_detail 人话），续跑后完成。"""

    def test_failed_then_resume_visible_via_api(self) -> None:
        hook = RecordingHook(fail_rows={"坏行标题"}, transient=False)
        app, service, hook = build_app(hook)
        data = build_csv_bytes(
            PRODUCT_HEADER,
            [valid_row("好行", "正文", "SKU-001"), valid_row("坏行标题", "正文", "SKU-002")],
        )
        with TestClient(app) as client:
            job_id = upload(client, file_name="p.csv", data=data).json()["data"]["job_id"]
            final = wait_done(client, job_id)
            assert final["status"] == "failed"
            assert "导入中断" in final["error_detail"]
            assert final["valid_rows"] == 1                      # 好行已入库（部分成功）

            # 「修复」后直接再触发同一任务：水位续跑（无重复创建）
            hook.fail_rows.clear()
            import asyncio

            async def resume() -> None:
                await service.process(job_id)

            asyncio.run(resume())
            resumed = client.get(f"{IMPORT_PATH}/{job_id}").json()["data"]
            assert resumed["status"] == "done"
            assert (resumed["valid_rows"], resumed["error_rows"]) == (2, 0)
            assert len(hook.items) == 2
