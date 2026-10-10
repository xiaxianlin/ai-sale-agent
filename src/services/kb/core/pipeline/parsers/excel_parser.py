# 知识产权归甲方所有（COM-IP-01）——AI 外贸销售智能体系统交付物。
# 许可与再分发须依据双方合同约定执行（COM-IP-02 乙方零使用权）。
"""Excel 解析器（FR-KB-010）：openpyxl 多 sheet 结构化解析。

每 sheet 产出 1 个 heading 块（“工作表：<名称>”）+ N 个完整表格块；
连续 ≥2 空行视为逻辑表边界，每张逻辑表整块发射、绝不按行数切割
（与 FR-KB-012 表格不跨块契约衔接）。

实现方案：docs/07-实现方案/KB/SR-KB-010-多格式文档解析.md §6.3、§7
"""
# 知识产权归甲方所有（COM-IP-01）

from __future__ import annotations

from typing import Any, Mapping

from .base import (
    BLOCK_TYPE_HEADING,
    BaseParser,
    Block,
    CorruptOrEncryptedError,
    UnsupportedFormatError,
    cfg_get,
)


def _load_openpyxl() -> Any:
    """惰性加载 openpyxl（import 守护）：未安装时抛带中文提示的不支持格式错误。"""
    try:
        import openpyxl  # noqa: F401 —— 仅探测可导入性
    except ImportError as exc:
        raise UnsupportedFormatError(
            "未安装 openpyxl，无法解析 Excel 文件（.xlsx/.xlsm）；请安装依赖：pip install openpyxl"
        ) from exc
    return openpyxl


class ExcelParser(BaseParser):
    """openpyxl：多 sheet 结构化。每 sheet 产出 1 个 heading 块 + N 个完整表格块。"""

    parser_name = "excel_parser"
    fmt = "excel"

    def __init__(self, cfg: Mapping[str, Any] | None = None) -> None:
        # 配置全部外置可覆盖（app_config 热更新口径，NFR-MAI-02；默认值见实现方案 §7）
        self.cfg: dict[str, Any] = {"parse.excel.max_cells": 500000}
        if cfg:
            self.cfg.update(cfg)

    def _do_parse(self, path: str) -> list[Block]:
        openpyxl = _load_openpyxl()
        try:
            # data_only=True 取公式计算值防信息丢失（§16.1 数据完整性）
            wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            raise CorruptOrEncryptedError(
                f"Excel 文件无法打开（可能为旧版 .xls 或已损坏）：{exc}"
            ) from exc

        blocks: list[Block] = []
        warnings: list[str] = []
        try:
            for ws in wb.worksheets:
                self._check_cells(ws, warnings)
                # 合并单元格值在左上角、余格为空串（read_only 模式下非锚点格为 None）
                rows = [
                    ["" if c is None else str(c).strip() for c in row]
                    for row in ws.iter_rows(values_only=True)
                ]
                rows = self._trim_blank(rows)  # 去首尾全空行/列
                if not rows:
                    continue  # 整 sheet 为空：不产块
                blocks.append(Block(type=BLOCK_TYPE_HEADING, heading_level=1, text=f"工作表：{ws.title}"))
                # 连续≥2 空行 = 逻辑表边界；每张逻辑表整块发射，绝不按行数切割
                for sub in self._split_by_blank_rows(rows):
                    blocks.append(self.emit_table(sub, {"sheet": ws.title}))
        finally:
            wb.close()  # read_only 句柄及时释放，防文件句柄泄漏
        if warnings:
            self._set_stats(warnings=warnings)
        return blocks

    def _check_cells(self, ws: Any, warnings: list[str]) -> None:
        """单 sheet 单元格数超 parse.excel.max_cells → 记告警供人工确认（防误传大文件拖垮流水线）。"""
        max_cells = int(self.cfg["parse.excel.max_cells"])
        if ws.max_row is None or ws.max_column is None:  # read_only 且缺 dimension 信息时跳过检查
            return
        cell_count = ws.max_row * ws.max_column
        if cell_count > max_cells:
            warnings.append(
                f"工作表 {ws.title} 单元格数 {cell_count} 超上限 {max_cells}，请人工确认是否误传大文件"
            )

    @staticmethod
    def _trim_blank(rows: list[list[str]]) -> list[list[str]]:
        """去首尾全空行/全空列；行宽不齐时右侧补空串规范化为矩形。"""
        if not rows:
            return []
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        # 首尾全空行
        while normalized and all(cell == "" for cell in normalized[0]):
            normalized.pop(0)
        while normalized and all(cell == "" for cell in normalized[-1]):
            normalized.pop()
        if not normalized:
            return []
        # 首尾全空列
        while all(row and row[0] == "" for row in normalized):
            normalized = [row[1:] for row in normalized]
        while all(row and row[-1] == "" for row in normalized):
            normalized = [row[:-1] for row in normalized]
        return normalized

    @staticmethod
    def _split_by_blank_rows(rows: list[list[str]], min_blank_run: int = 2) -> list[list[list[str]]]:
        """以「连续 >=min_blank_run 空行」为逻辑表边界切分；单空行保留在表内。

        返回若干非空逻辑表；每个逻辑表内部再去掉首尾残留空行。
        """
        groups: list[list[list[str]]] = []
        current: list[list[str]] = []
        blank_run = 0
        for row in rows:
            if all(cell == "" for cell in row):
                blank_run += 1
                if blank_run >= min_blank_run:
                    if current:
                        groups.append(current)  # 边界确认：当前逻辑表结束
                        current = []
                elif current:
                    current.append(row)  # 单空行属表内结构，保留（如行距分隔）
                continue
            blank_run = 0
            current.append(row)
        if current:
            groups.append(current)
        # 每张逻辑表去首尾空行（单个边界空行残留），保证表块干净
        trimmed: list[list[list[str]]] = []
        for group in groups:
            start, end = 0, len(group)
            while start < end and all(cell == "" for cell in group[start]):
                start += 1
            while end > start and all(cell == "" for cell in group[end - 1]):
                end -= 1
            if end > start:
                trimmed.append(group[start:end])
        return trimmed


def parse_excel_config(cfg: Mapping[str, Any] | None) -> dict[str, Any]:
    """辅助函数：合并 Excel 解析配置默认值（供编排层/测试复用）。"""
    merged = {"parse.excel.max_cells": int(cfg_get(cfg, "parse.excel.max_cells", 500000))}
    return merged
